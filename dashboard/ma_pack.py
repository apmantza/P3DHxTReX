"""Deal pack sheets of the M&A workbook: Valuation, Balance sheet, P&L, MREL scenario, Decisions, and Sources. NBG and the first target (BIL) only.
write_sheets() is called by ma_workbook.build_model_sheets when the market data holds a pack (ma_pack_data.load). Labels and notes follow ASD-STE100."""

from __future__ import annotations

from collections import Counter
from typing import Any

import ma_pack_pnl
import ma_workbook as W
from openpyxl.styles import Alignment, Font

from ma_pack_base import BLUE, GREEN, HEAD, HEAD_FONT, INPUT_FILL, PNL_SHEET, SECTION, Sheet

VAL, BS, MREL, DEC, SRC = "Valuation", "Balance sheet", "MREL scenario", "Decisions", "Sources"
NEW_SHEETS = [VAL, BS, PNL_SHEET, MREL, DEC, SRC]
BASIS_NAMES = '"Tangible equity","CET1","Average of the two"'


def q(name: str) -> str:
    """A sheet name for a formula."""
    return f"'{name}'" if any(ch in name for ch in " &") else name


# ================================================================================================ Valuation
def write_valuation(ws, banks: list, deals: list, g: dict, market: dict, pack: dict, pnl_rows: dict[str, int]) -> dict[str, int]:
    S = Sheet(ws, (74, 15, 15, 15, 15, 15, 15, 100))
    name = deals[0]["short"]
    ws["A1"] = f"Valuation of {name}: how the P/TBV and P/CET1 comps are derived. EUR billion unless stated. Green: data. Blue on yellow: input. Black: formula."
    ws["A1"].font = Font(bold=True)
    ws["A2"] = f"{name} has no listed shares. The model values it with the median multiple of its nearest listed peers. The sections run from the peer set to the value."
    S.r = 4
    S.head(["1. Definitions", "", "", "", "", "", "", ""])
    for line in (
        "P/TBV = price to book ratio x book value per share / tangible book value per share. Bloomberg gives the per-share values in EUR and the price in the currency of the stock.",
        "The price to book ratio is the same in any currency, so the model uses it. RoTE = EPS of the last 12 months / tangible book value per share.",
        "P/CET1 = market cap in EUR / CET1 of the newest KM1 filing in P3DH. RoCET1 = P/CET1 x EPS / price in EUR.",
        "The nearest peers are the k banks whose return is closest to the return of the target. The model uses RoTE for P/TBV and RoCET1 for P/CET1.",
        "The comp is the median multiple of the nearest peers. A tie at the k-th distance puts all tied banks in the set.",
        "Gordon model: P/TBV = (RoTE - g) / (cost of equity - g). The cost of equity of a peer is g + (RoTE - g) / P/TBV.",
        f"The model scales the peer value by (cost of equity of the peers - g) / (cost of equity of {name} - g). {name} is smaller and unlisted, so its cost of equity is higher.",
    ):
        S.text(line)
    S.skip()

    # ------------------------------------------------------------------ funnel
    S.head(["2. Peer funnel: from the Bloomberg universe to the peer set", "Banks left", "Banks removed", "", "", "", "", "Source and note"])
    funnel = pack.get("funnel") or []
    if not funnel:
        S.text("The funnel is not available: the peer data is missing.")
    n_rows = len(market["peers"]["rows"])
    for i, (label, n) in enumerate(funnel):
        S.row(f"fun{i}", label, [n, None if i == 0 else f"=B{S.r - 1}-B{S.r}"], "m", "Counts: sheet Peers and ma_market.peers()." if i == 0 else "", ["data"])
    if funnel:
        first, last = W.PEER_TOP, W.PEER_TOP + n_rows - 1
        S.row("fun_check", "Banks on the sheet Peers (check: equal to the last step; 1 = yes)", [f"=COUNT(Peers!D{first}:D{last})", f"=IF(B{S.r}=B{S.rows['fun' + str(len(funnel) - 1)]},1,0)"], ["m", "flag"])
        by = pack.get("removed_by_reason") or {}
        scale = funnel[2][1] - funnel[3][1] if len(funnel) > 3 else 0
        for label, n, note in (("Removed in step 4 for a scale mismatch (TREA of P3DH / RWA of Bloomberg below 0.5 or above 2)", scale, "The list of the removed banks is at the foot of the sheet Peers."),
                               ("Removed in step 6 as a subsidiary (Bloomberg gives the bank a parent in a group)", by.get("subsidiary", 0), ""),
                               ("Removed in step 6 as a specialist (leasing, commercial finance, brokerage, investment management, mortgage, special purpose)", by.get("specialist", 0), ""),
                               ("Removed in step 6 as a cooperative (Bloomberg subgroup, legal form, or name)", by.get("cooperative", 0), "")):
            S.row(None, label, [None, n], "m", note, ["data", "data"])
    S.skip()

    # ------------------------------------------------------------------ inputs of the target
    pcol = W.peer_cols(0)
    S.head([f"3. {name}: the inputs of the comparison", name, "", "", "", "", "", "Source and note"])
    S.row("year", "Year of the analysis (Calc N15)", ['=CHOOSE(Calc!$N$15,2026,2027,2028,2029)'], "int", "Lever N15 of Calc.")
    S.row("profit", "Profit of the analysis year (EUR bn)", ["=Calc!E9"], "bn", f"{name} has no consensus: Calc row 9 holds the trailing profit from continuing operations (Bloomberg). The driver-based profit is on the sheet P&L.")
    S.row("tbv", "Tangible equity (EUR bn)", ["=Calc!E30"], "bn", "P3DH CC1: common equity before adjustments less intangible assets.")
    S.row("cet1", "CET1 (EUR bn)", ["=Calc!E20"], "bn", "P3DH KM1.")
    S.row("rote", "RoTE = profit / tangible equity", ["=Calc!E161"], "pct")
    S.row("rocet1", "RoCET1 = profit / CET1", ["=Calc!E162"], "pct")
    S.row("k", "Nearest peers used (k, Calc N17)", ["=Calc!N17"], "int")
    S.row("basis", "Valuation basis (Calc N18)", [f"=CHOOSE(Calc!N18,{BASIS_NAMES})"], "text")
    S.skip()

    # ------------------------------------------------------------------ the comps
    S.head(["4. The comps: median multiple of the nearest peers", "On tangible equity (P/TBV)", "On CET1 (P/CET1)", "", "", "", "", "Source and note"])
    S.row("c_ret", "Return of the target used to pick the peers (RoTE; RoCET1)", ["=Calc!E161", "=Calc!E162"], "pct", "Calc rows 161 and 162.")
    S.row("c_dist", "Distance of the k-th nearest peer (points of return)", [f"=Peers!{pcol['d_tbv']}6", f"=Peers!{pcol['d_cet1']}6"], "pct", "Sheet Peers, row 6.")
    S.row("c_n", "Peers in the set", [f"=Peers!{pcol['m_tbv']}8", f"=Peers!{pcol['m_cet1']}8"], "int", "More than k when peers tie at the k-th distance.")
    S.row("c_med", "Median multiple of the nearest peers (x)", ["=Calc!E163", "=Calc!E164"], "x", "Calc rows 163 and 164. This is the comp.", None, True)
    S.row("c_coe_p", "Cost of equity that the nearest peers imply", ["=Calc!E226", "=Calc!E227"], "pct", "Calc rows 226 and 227: g + (return - g) / multiple, median.")
    S.row("c_coe_t", f"Cost of equity of {name} (Calc N24)", ["=Calc!E225", "=Calc!E225"], "pct", "0% means: use the cost of equity of the peers (no factor).")
    S.row("c_g", "Growth rate g (Calc N23)", ["=Calc!N23", "=Calc!N23"], "pct")
    S.row("c_fac", "Factor = (cost of equity of the peers - g) / (cost of equity of the target - g)", ["=Calc!E228", "=Calc!E229"], "f3", "Calc rows 228 and 229.")
    S.row("c_adj", "Multiple after the factor (x)", [f"=B{S.rows['c_med']}*B{S.rows['c_fac']}", f"=C{S.rows['c_med']}*C{S.rows['c_fac']}"], "x", "", None, True)
    S.row("c_val", "Value of the target (EUR bn)", ["=Calc!E165", "=Calc!E166"], "bn", "Multiple x tangible equity (P/TBV) or x CET1 (P/CET1), after the factor.")
    S.row("c_used", "Value used in the model (EUR bn): basis in Calc N18", ["=Calc!E167"], "bn", "Calc row 167. It is the market cap of the target (Calc row 32).", None, True)
    S.row("c_eff", "Effect of the cost of equity factor on the value (EUR bn)", ["=Calc!E231"], "bn", "Calc row 231.")
    S.skip()

    # ------------------------------------------------------------------ football field
    S.head([f"5. Value of {name} (100%) by method", "Value (EUR bn)", "P/TBV (x)", "P/CET1 (x)", "P/E (x)", "Price for the stake (EUR bn)", "Against the value used", "Source and note"])
    price = lambda r: f"=Calc!$H$4*B{r}*(1+Calc!$E$37)"  # noqa: E731

    def method(key: str, label: str, formula: str, note: str, bold: bool = False) -> None:
        r = S.r
        S.row(key, label, [formula, f"=B{r}/Calc!$E$30", f"=B{r}/Calc!$E$20", f"=IF(Calc!$E$9>0,B{r}/Calc!$E$9,0)", price(r), f"=B{r}/Calc!$E$167-1"], ["bn", "x", "x", "x", "bn", "pct"], note, None, bold)

    gor = lambda rote: f"=IF(Calc!$N$24>Calc!$N$23,Calc!$E$30*({rote}-Calc!$N$23)/(Calc!$N$24-Calc!$N$23),0)"  # noqa: E731
    method("f_tbv", "P/TBV of the nearest peers, before the cost of equity factor", "=Calc!E163*Calc!E30", "Comp x tangible equity.")
    method("f_cet1", "P/CET1 of the nearest peers, before the cost of equity factor", "=Calc!E164*Calc!E20", "Comp x CET1.")
    method("f_avg", "Average of the two, before the cost of equity factor", f"=AVERAGE(B{S.rows['f_tbv']},B{S.rows['f_cet1']})", "")
    method("f_tbv2", "P/TBV after the cost of equity factor", "=Calc!E165", "Calc row 165.")
    method("f_cet12", "P/CET1 after the cost of equity factor", "=Calc!E166", "Calc row 166.")
    method("f_used", "Value used in the model (basis in Calc N18)", "=Calc!E167", "Calc row 167. The checks compare this row with Calc rows 32 and 33.", True)
    method("f_gord", "Gordon model on the trailing RoTE, at the cost of equity of the target", gor("Calc!$E$161"), "Tangible equity x (RoTE - g) / (cost of equity - g). It needs N24 above g.")
    pnl = q(PNL_SHEET)
    rr = pnl_rows["rote_fwd"]
    method("f_gordf", "Gordon model on the RoTE of the driver-based profit of the analysis year", gor(f"CHOOSE(Calc!$N$15,{pnl}!C{rr},{pnl}!D{rr},{pnl}!E{rr},{pnl}!F{rr})"),
           "The RoTE comes from the sheet P&L (switch there: drivers or flat).")
    method("f_pe", f"P/E of NBG x profit of {name} (reference only)", "=Calc!E9*Calc!N3", "Calc N3 is the P/E of NBG on the same basis. A listed bank of a different country and size has another P/E.")
    r = S.r
    S.row("f_deal", f"Deal: price for the stake, scaled to 100% of {name}", ["=IF(Calc!H4>0,Calc!E38/Calc!H4,0)", f"=B{r}/Calc!$E$30", f"=B{r}/Calc!$E$20", f"=IF(Calc!$E$9>0,B{r}/Calc!$E$9,0)", "=Calc!E38", f"=B{r}/Calc!$E$167-1"],
          ["bn", "x", "x", "x", "bn", "pct"], "Calc row 38 / stake. The premium (Calc E37) lifts the price above the value used.", None, True)
    S.skip()

    # ------------------------------------------------------------------ cost of equity grid, live
    S.head([f"6. Cost of equity of {name}: live table (g, k, and the peers as in section 4)", "Cost of equity", "Factor, P/TBV", "Value, P/TBV (EUR bn)", "Factor, P/CET1", "Value, P/CET1 (EUR bn)", "Average (EUR bn)", "Source and note"])
    for coe in (0.09, 0.10, 0.11, 0.12, 0.13):
        r = S.r
        fa = f"IF(AND(B{r}>Calc!$N$23,Calc!$E$226>Calc!$N$23),(Calc!$E$226-Calc!$N$23)/(B{r}-Calc!$N$23),1)"
        fb = f"IF(AND(B{r}>Calc!$N$23,Calc!$E$227>Calc!$N$23),(Calc!$E$227-Calc!$N$23)/(B{r}-Calc!$N$23),1)"
        S.row(None, "Cost of equity of the target (input of this row)", [coe, f"={fa}", f"=Calc!$E$163*Calc!$E$30*C{r}", f"={fb}", f"=Calc!$E$164*Calc!$E$20*E{r}", f"=AVERAGE(D{r},F{r})"],
              ["pct", "f3", "bn", "f3", "bn", "bn"], "Calc N24 is the lever of the model. This table does not change it." if coe == 0.09 else "", ["input"])
    S.skip()

    # ------------------------------------------------------------------ static grids
    kg, gg = pack.get("k_grid") or [], pack.get("g_grid") or []
    S.head(["7. Number of nearest peers k: STATIC table, values at the build date. It does not follow the levers.", "k", "Peers, P/TBV", "Median P/TBV (x)", "Peers, P/CET1", "Median P/CET1 (x)", "Average value (EUR bn)", "Source and note"])
    if not kg:
        S.text("Not available: the inputs of the target are missing.")
    for i, x in enumerate(kg):
        S.row(None, "Before the cost of equity factor", [x["k"], x["n_tbv"], x["ptbv"], x["n_cet1"], x["pcet1"], (x["value_tbv"] + x["value_cet1"]) / 2], ["int", "int", "x", "int", "x", "bn"],
              "Python calculation with the same rules as the sheet Peers (ties included)." if i == 0 else "", ["data"] * 6)
    S.skip()
    S.head(["8. Growth rate g: STATIC table, k = 20 and cost of equity of the target 10%", "g", "Cost of equity of the peers, P/TBV", "Factor, P/TBV", "Cost of equity of the peers, P/CET1", "Factor, P/CET1", "Average value (EUR bn)", "Source and note"])
    if not gg:
        S.text("Not available: the inputs of the target are missing.")
    for i, x in enumerate(gg):
        S.row(None, "After the cost of equity factor", [x["g"], x["coe_tbv"], x["factor_tbv"], x["coe_cet1"], x["factor_cet1"], (x["value_tbv"] + x["value_cet1"]) / 2], ["pct", "pct", "f3", "pct", "f3", "bn"],
              "A higher g lowers the cost of equity that the peers imply and lowers the factor." if i == 0 else "", ["data"] * 6)
    S.skip()

    # ------------------------------------------------------------------ Gordon check
    gd = pack.get("gordon") or {}
    S.head(["9. Check: cost of equity that the peers imply by three routes (STATIC)", "Peers", "Lower quartile", "Median", "Upper quartile", "", "", "Source and note"])
    if not gd:
        S.text("Not available.")
    for key, label, note in (("trailing", "Trailing RoTE and the market multiple", "The route of the model (sheet Peers)."),
                             ("consensus", "RoTE of the consensus for the first fiscal year and the market multiple", "A forward return makes the peers look cheaper: the implied cost of equity is lower."),
                             ("target", "Consensus RoTE and the multiple of the analyst target price", "Shows what the analysts expect the market to pay.")):
        v = gd.get(key)
        if v and v.get("n"):
            S.row(None, label, [v["n"], v["q1"], v["median"], v["q3"]], ["int", "pct", "pct", "pct"], note, ["data"] * 4)
    S.skip()
    rows = market["peers"]["rows"]
    by_country = Counter(p["country"] for p in rows)
    S.text(f"Peers by country: {', '.join(f'{c} {n}' for c, n in sorted(by_country.items()))}." if rows else "No peers.")
    ws.freeze_panes = "B4"
    return S.rows


# ================================================================================================ Balance sheet
def write_balance(ws, banks: list, deals: list, market: dict, pack: dict) -> dict[str, int]:
    S = Sheet(ws, (74, 14, 14, 14, 16, 14, 14, 100))
    name = deals[0]["short"]
    ws["A1"] = f"Balance sheet and capital metrics of NBG and {name}. EUR billion unless stated. Green: data. Black: formula."
    ws["A1"].font = Font(bold=True)
    ws["A2"] = ("The simple sum adds the two banks without any effect of the deal. The pro forma column is deal 1 on the sheet Calc: price, goodwill, minorities, instruments, and cash. "
                "A ratio in the simple sum is the sum of the amounts divided by the sum of the bases.")
    S.r = 4
    cols = ["", "NBG", name, "Simple sum", "Pro forma (deal 1)", "", "", "Source and note"]
    r_ = lambda key: S.rows[key]  # noqa: E731

    def amount(key: str, label: str, calc_row: int, fmt: str = "bn", note: str = "", pro: bool = True, bold: bool = False) -> None:
        r = S.r
        S.row(key, label, [f"=Calc!D{calc_row}", f"=Calc!E{calc_row}", f"=B{r}+C{r}", f"=Calc!H{calc_row}" if pro else None], fmt, note, ["data", "data", None, None], bold)

    def ratio(key: str, label: str, calc_row: int, num: list[str], den: list[str], fmt: str = "pct", note: str = "") -> None:
        r = S.r
        simple = f"=(B{r_(num[0])}+C{r_(num[0])}{''.join(f'+B{r_(n)}+C{r_(n)}' for n in num[1:])})/(B{r_(den[0])}+C{r_(den[0])})" if num and den else None
        S.row(key, label, [f"=Calc!D{calc_row}", f"=Calc!E{calc_row}", simple, f"=Calc!H{calc_row}"], fmt, note, ["data", "data", None, None])

    S.head(["A. Capital"] + cols[1:])
    amount("assets", "Total assets", 7, note="P3DH. The pro forma adds the goodwill and the change of cash and removes the intangible assets of the target.")
    amount("trea", "Total risk exposure amount (TREA)", 8, note="P3DH KM1 row 4.")
    amount("cet1", "CET1", 20, note="P3DH KM1 row 1. The pro forma subtracts the goodwill, the cash paid, and the minorities (Calc rows 73 to 88).")
    ratio("cet1_r", "CET1 ratio", 21, ["cet1"], ["trea"])
    amount("at1", "AT1", 22, note="The pro forma counts the AT1 of the target at the recognition of CRR Article 87 (Calc N19 and row 233).")
    amount("t2", "Tier 2", 23)
    amount("t1", "Tier 1", 108)
    r = S.r
    S.row("tc", "Total capital", [f"=Calc!D20+Calc!D22+Calc!D23", f"=Calc!E20+Calc!E22+Calc!E23", f"=B{r}+C{r}", "=Calc!H20+Calc!H22+Calc!H23"], "bn", "", ["data", "data"])
    amount("tbv", "Tangible equity (common equity before adjustments less intangible assets)", 30, note="P3DH CC1.")
    r = S.r
    S.row("gw", "New goodwill of the deal", [None, None, None, "=Calc!H81"], "bn", "Calc row 81. A bargain purchase has no goodwill.", None)
    S.row("req_cet1", "CET1 requirement (Pillar 1 + P2R + combined buffer; the trigger of the MDA)", ["=Calc!D100", "=Calc!E100", None, "=Calc!H100"], "pct2", "P3DH KM1 and CBR. The pro forma weights the requirements by the RWAs.", ["data", "data"])
    S.row("req_tc", "Total capital requirement", ["=Calc!D102", "=Calc!E102", None, "=Calc!H102"], "pct2", "", ["data", "data"])
    S.row("p2r", "Pillar 2 requirement (P2R)", ["=Calc!D92", "=Calc!E92", None, "=Calc!H92"], "pct2", "P3DH KM1 row 7d.", ["data", "data"])
    S.row("cbr", "Combined buffer requirement", ["=Calc!D99", "=Calc!E99", None, "=Calc!H99"], "pct2", "P3DH KM1 row 11.", ["data", "data"])
    S.row("mda", "Distance to the MDA trigger (bp of TREA)", ["=Calc!D106", "=Calc!E106", None, "=Calc!H106"], "bp", "The smallest headroom of CET1, Tier 1, and total capital.", ["data", "data"])
    S.row("bind", "Binding layer", ["=Calc!D107", "=Calc!E107", None, "=Calc!H107"], "text", "", ["data", "data"])
    S.skip()

    S.head(["B. Leverage and liquidity"] + cols[1:])
    amount("lev_exp", "Leverage ratio exposure measure", 109, note="P3DH KM1 row 13.")
    ratio("lev", "Leverage ratio (Tier 1 / exposure)", 110, ["t1"], ["lev_exp"])
    S.row("lev_req", "Leverage requirement", ["=Calc!D111", "=Calc!E111", None, "=Calc!H111"], "pct2", "P3DH KM1 row 14e. The model uses 3% when the bank files none.", ["data", "data"])
    S.row("lev_hd", "Leverage headroom (bp)", ["=Calc!D112", "=Calc!E112", None, "=Calc!H112"], "bp", "", ["data", "data"])
    amount("hqla", "Liquidity buffer (HQLA)", 113, note="P3DH KM1 row 15.")
    amount("outflow", "Net cash outflows", 114, note="P3DH KM1 row 16.")
    ratio("lcr", "Liquidity coverage ratio (LCR)", 115, ["hqla"], ["outflow"])
    amount("asf", "Available stable funding", 116, note="P3DH KM1 row 18.")
    amount("rsf", "Required stable funding", 117, note="P3DH KM1 row 19.")
    ratio("nsfr", "Net stable funding ratio (NSFR)", 118, ["asf"], ["rsf"])
    S.row("enc", "Encumbered assets / total assets (AE1)", ["=Calc!D127", "=Calc!E127", None, "=Calc!H127"], "pct", "P3DH AE1.", ["data", "data"])
    S.skip()

    S.head(["C. MREL (detail on the sheet MREL scenario)"] + cols[1:])
    amount("sen", "Eligible liabilities (senior)", 24, note="P3DH KM2: own funds and eligible liabilities less total capital. The pro forma counts the senior debt of the target at Calc N26.")
    amount("mrel", "MREL resources", 25)
    ratio("mrel_r", "MREL ratio (% of TREA)", 26, ["mrel"], ["trea"])
    S.row("mrel_q", "MREL requirement with the combined buffer (% of TREA)", ["=Calc!D120", "=Calc!E120", None, "=Calc!H120"], "pct2", "P3DH KM2 row EU-7 and the combined buffer.", ["data", "data"])
    S.row("mrel_h", "MREL headroom (bp of TREA)", ["=Calc!D121", "=Calc!E121", None, "=Calc!H121"], "bp", "", ["data", "data"])
    W_ = W.p3dh_rows()
    S.row("mrel_tem", "MREL requirement as % of the leverage exposure (KM2 row EU-9)", [f"=P3DH!D{W_['km2_req_tem']}", f"=P3DH!E{W_['km2_req_tem']}"], "pct2", "The requirement on the exposure can bind before the requirement on TREA.", ["data", "data"])
    S.skip()

    S.head(["D. Risk-weighted assets by type (OV1)"] + cols[1:])
    types = [("rwa_credit", "Credit risk"), ("rwa_ccr", "Counterparty credit risk"), ("rwa_cva", "Credit valuation adjustment"), ("rwa_securitisation", "Securitisation"),
             ("rwa_market", "Market risk"), ("rwa_operational", "Operational risk")]
    first = S.r
    for key, label in types:
        r = S.r
        S.row("t_" + key, label, [f"=P3DH!D{W_[key]}", f"=P3DH!E{W_[key]}", f"=B{r}+C{r}", f"=D{r}*Calc!$H$8/(Calc!$D$8+Calc!$E$8)"], "bn",
              "P3DH OV1. The pro forma scales the simple sum to the TREA of Calc (levers D65 and E65)." if key == "rwa_credit" else "", ["data", "data"])
    r = S.r
    S.row("t_sum", "Sum of the types", [f"=SUM(B{first}:B{r - 1})", f"=SUM(C{first}:C{r - 1})", f"=SUM(D{first}:D{r - 1})", f"=SUM(E{first}:E{r - 1})"], "bn", "The sum can differ a little from TREA: OV1 has other lines (settlement, large exposures).", None, True)
    S.row("t_op", "Share of operational risk in TREA", [f"=IF(B{r_('trea')}>0,B{r_('t_rwa_operational')}/B{r_('trea')},0)", f"=IF(C{r_('trea')}>0,C{r_('t_rwa_operational')}/C{r_('trea')},0)",
                                                      f"=IF(D{r_('trea')}>0,D{r_('t_rwa_operational')}/D{r_('trea')},0)", f"=IF(E{r_('trea')}>0,E{r_('t_rwa_operational')}/E{r_('trea')},0)"], "pct")
    S.skip()

    S.head(["E. Asset quality (CR1)"] + cols[1:])
    cn, ct = pack.get("cr1_nbg") or {}, pack.get("cr1_tgt") or {}
    S.row("gross", "Gross loans and advances (EUR bn)", [cn.get("gross", 0) / 1000 if cn else None, ct.get("gross", 0) / 1000 if ct else None, f"=B{S.r}+C{S.r}"], "bn",
          f"P3DH CR1 sheet a. Dates: NBG {cn.get('date', 'n/a')}, {name} {ct.get('date', 'n/a')}.", ["data", "data"])
    S.row("npe", "Non-performing exposure (EUR bn)", [cn.get("npe", 0) / 1000 if cn else None, ct.get("npe", 0) / 1000 if ct else None, f"=B{S.r}+C{S.r}"], "bn", "Gross loans x non-performing ratio.", ["data", "data"])
    S.row("npl", "Non-performing ratio", [f"=IF(B{r_('gross')}>0,B{r_('npe')}/B{r_('gross')},0)", f"=IF(C{r_('gross')}>0,C{r_('npe')}/C{r_('gross')},0)", f"=IF(D{r_('gross')}>0,D{r_('npe')}/D{r_('gross')},0)"], "pct2")
    S.row("cov", "Coverage of non-performing exposure: accumulated impairment / exposure", [(cn.get("impairment_npe", 0) / cn["npe"]) if cn.get("npe") else None, (ct.get("impairment_npe", 0) / ct["npe"]) if ct.get("npe") else None], "pct",
          "P3DH CR1 sheet d. Empty when the filing has no impairment of the non-performing loans.", ["data", "data"])
    S.skip()
    mixes = [cn.get("mix") or {}, ct.get("mix") or {}]
    names = sorted(set(mixes[0]) | set(mixes[1]), key=lambda k: -(mixes[1].get(k, 0) + mixes[0].get(k, 0)))
    S.head(["F. Loans by counterparty (% of the gross loans, CR1)"] + cols[1:])
    if not names:
        S.text("Not available: CR1 has no counterparty split.")
    for i, nm in enumerate(names):
        r = S.r
        S.row(None, nm, [mixes[0].get(nm, 0) / 100, mixes[1].get(nm, 0) / 100, f"=IF(B{r_('gross')}+C{r_('gross')}>0,(B{r}*B{r_('gross')}+C{r}*C{r_('gross')})/(B{r_('gross')}+C{r_('gross')}),0)"], "pct",
              "P3DH CR1. The simple sum weights the shares by the gross loans." if i == 0 else "", ["data", "data"])
    S.skip()

    S.head(["G. Interest rate risk in the banking book (IRRBB1)"] + cols[1:])
    for i, (code, label) in enumerate(W.ma_data.IRRBB_ROWS.items()):
        S.row(f"eve{code}", f"Change of economic value of equity, {label.lower()}, % of Tier 1", [f"=Calc!D{138 + i}", f"=Calc!E{138 + i}", None, f"=Calc!H{138 + i}"], "pct",
              "P3DH K_68.00. The supervisory outlier limit is 15% of Tier 1." if i == 0 else "", ["data", "data"])
    nn, nt = pack.get("irrbb_nii_nbg") or {}, pack.get("irrbb_nii_tgt") or {}
    for code, label in (("0010", "+200 bp"), ("0020", "-200 bp")):
        a, b = nn.get(code), nt.get(code)
        r = S.r
        S.row("nii" + code, f"Change of net interest income, parallel shift {label} (EUR m)", [a[0] / 1e6 if a else None, b[0] / 1e6 if b else None, f"=B{r}+C{r}"], "m1",
              f"P3DH K_68.00. Dates: NBG {a[1] if a else 'n/a'}, {name} {b[1] if b else 'n/a'}." if code == "0010" else "", ["data", "data"])
    S.skip()

    S.head(["H. Balance sheet lines (CC2, accounting scope; EUR bn)"] + cols[1:])
    c2n, c2t = pack.get("cc2_nbg") or {}, pack.get("cc2_tgt") or {}
    if not c2n and not c2t:
        S.text("Not available: CC2 does not tie to total assets.")
    S.text(f"Dates of the filings: NBG {c2n.get('date', 'n/a')}, {name} {c2t.get('date', 'n/a')}. The lines come from keywords of the CC2 labels. Each side ties to total assets within 2%.")
    asset_names = (("cash", "Cash and central bank balances"), ("interbank", "Loans and advances to credit institutions"), ("loans", "Loans and advances to customers"),
                   ("securities", "Securities and other financial assets"), ("derivatives", "Derivatives and hedging"), ("intangibles", "Intangible assets and goodwill"), ("other", "Other assets"))
    liab_names = (("interbank", "Amounts due to credit institutions"), ("deposits", "Amounts due to customers (deposits)"), ("debt", "Debt securities issued"),
                  ("sub", "Subordinated liabilities and other borrowed funds"), ("derivatives", "Derivatives and other financial liabilities"), ("other", "Other liabilities"))

    def cc2_rows(prefix: str, group: str, items: tuple) -> None:
        for key, label in items:
            r = S.r
            S.row(f"{prefix}_{key}", label, [c2n[group][key] / 1000 if c2n else None, c2t[group][key] / 1000 if c2t else None, f"=B{r}+C{r}"], "bn", "", ["data", "data"])

    S.section("Assets")
    cc2_rows("a", "assets", asset_names)
    a0, a1 = S.rows["a_cash"], S.rows["a_other"]
    S.row("a_tot", "Total assets of the lines", [f"=SUM(B{a0}:B{a1})", f"=SUM(C{a0}:C{a1})", f"=SUM(D{a0}:D{a1})"], "bn", "", None, True)
    S.row("a_chk", "Difference to total assets of Calc (EUR bn); a small value ties", [f"=B{r_('a_tot')}-Calc!D7", f"=C{r_('a_tot')}-Calc!E7"], "f3", "Dates can differ: CC2 and KM1 can have different reference dates.")
    S.section("Liabilities and equity")
    cc2_rows("l", "liabilities", liab_names)
    r = S.r
    S.row("l_eq", "Equity (accounting scope)", [c2n["equity"] / 1000 if c2n else None, c2t["equity"] / 1000 if c2t else None, f"=B{r}+C{r}"], "bn", "", ["data", "data"])
    l0 = S.rows["l_interbank"]
    S.row("l_tot", "Total liabilities and equity of the lines", [f"=SUM(B{l0}:B{r})", f"=SUM(C{l0}:C{r})", f"=SUM(D{l0}:D{r})"], "bn", "", None, True)
    S.row("l_chk", "Difference to the assets of the lines (EUR bn); a small value ties", [f"=B{r_('l_tot')}-B{r_('a_tot')}", f"=C{r_('l_tot')}-C{r_('a_tot')}"], "f3")
    S.row("l_ld", "Loans to customers / deposits", [f"=IF(B{r_('l_deposits')}>0,B{r_('a_loans')}/B{r_('l_deposits')},0)", f"=IF(C{r_('l_deposits')}>0,C{r_('a_loans')}/C{r_('l_deposits')},0)", f"=IF(D{r_('l_deposits')}>0,D{r_('a_loans')}/D{r_('l_deposits')},0)"], "pct")
    S.skip()

    S.head(["I. Bonds outstanding by rank (Bloomberg, EUR bonds only; EUR bn)"] + cols[1:])
    bn_, bt_ = pack.get("bonds_nbg") or {}, pack.get("bonds_tgt") or {}
    for key, label in (("senior_preferred", "Senior preferred"), ("senior_non_preferred", "Senior non-preferred"), ("other_senior", "Other ranks: unsecured, secured, or not classed by Bloomberg"), ("tier2", "Tier 2"), ("at1", "AT1")):
        r = S.r
        S.row("b_" + key, label, [bn_.get(key), bt_.get(key), f"=B{r}+C{r}"], "bn",
              (f"Bloomberg bond census. Bonds in other currencies are not summed: NBG {bn_.get('non_eur_count', 'n/a')}, {name} {bt_.get('non_eur_count', 'n/a')}. "
               "A sum that mixes currencies is wrong: the model uses KM2 for the senior stack.") if key == "senior_preferred" else "", ["data", "data"])
    S.skip()

    S.head(["J. Reference dates and data quality"] + cols[1:])
    S.row("ref", "Reference quarter of P3DH (KM1 filing)", ["=P3DH!D3", "=P3DH!E3"], "text", "Sheet P3DH, row 3. The model needs the same quarter for both banks: Calc row 149 checks it.", ["data", "data"])
    S.row("chk", "Checks of the sheet Calc passed", ["=Calc!D151"], "text", "Calc row 151.")
    for m in pack.get("missing") or []:
        S.text(f"Missing data: {m}")
    ws.freeze_panes = "B5"
    return S.rows


# ================================================================================================ Decisions
def write_decisions(ws, banks: list, deals: list, market: dict, pack: dict, rows: dict[str, dict[str, int]]) -> None:
    name = deals[0]["short"]
    widths = (5, 44, 62, 18, 14, 16, 50, 40, 52, 10)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws["A1"] = "Decisions to make: the choices behind the numbers, with a live link to the lever and to its effect. The model keeps the default shown. Green: link."
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "A decision is open until someone with authority takes it. The column Effect shows how much the number moves now."
    heads = ["No.", "Decision", "Options", "Lever", "Value now", "Effect now", "Effect measured as", "Default in the model", "Who decides, and note", "Status"]
    for j, h in enumerate(heads):
        c = ws.cell(row=4, column=1 + j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ml, pn = rows[MREL], rows[PNL_SHEET]
    R = W.merger_start(0)
    mg = lambda key, col: f"Merger!{col}{R + W.MOFF[key]}"  # noqa: E731
    by_year = lambda fn: "=CHOOSE(Calc!$N$15," + ",".join(fn(c) for c in W.YC) + ")"  # noqa: E731
    by_pnl = lambda key: by_year(lambda c: f"({q(PNL_SHEET)}!{chr(ord(c) - 1)}{pn['np_drv']}-{q(PNL_SHEET)}!{chr(ord(c) - 1)}{pn['np_flat']})")  # noqa: E731
    items = [
        ("Resolution strategy of the group: single point of entry (SPE) or multiple points of entry (MPE)",
         "1 = SPE: NBG is the resolution entity and BIL is a subsidiary. 2 = MPE: BIL stays its own resolution entity.", f"{MREL} B4", f"={q(MREL)}!B4", "int",
         f"={q(MREL)}!D{ml['sel_need']}", "bn", "Issuance need at the requirements for the strategy selected (EUR bn)",
         "SPE. The precedent of Eurobank and Hellenic Bank points to it.", "The SRB decides. NBG can argue its case in the MREL decision. The law does not give a default."),
        ("Stake bought", "Any share from 50% to 100%. The State of Luxembourg holds 10% and does not sell.", "Calc H4", "=Calc!H4", "pct", "=Calc!E38", "bn", "Price for the stake (EUR bn)",
         "90%", "NBG and the sellers. A stake below 100% leaves minorities: Calc rows 84 and 85 treat them."),
        ("Premium over the peer value", "No premium, or a premium that the sellers ask for.", "Calc E37", "=Calc!E37", "pct", "=Calc!E38-Calc!H4*Calc!E32", "bn", "Premium in the price for the stake (EUR bn)",
         "0%", "NBG and the sellers. A premium over the value raises the goodwill, or lowers the bargain gain, and lowers CET1 one for one."),
        ("Cash share of the price", "100% cash, 100% shares, or a mix. Cash costs the deposit rate on the cash paid.", "Calc H2", "=Calc!H2", "pct", "=Calc!H10", "bn",
         "Income lost on the cash per year, after tax (EUR bn)", "90%", "NBG. The deposit rate path is in Market row 87."),
        ("Recognition of the AT1 and Tier 2 of BIL in the own funds of the group", "100% = the amount of CRR Article 87. A lower share tests a stricter reading by the supervisor.", "Calc N19", "=Calc!N19", "pct",
         "=Calc!N19*Calc!E233*(Calc!E22+Calc!E23)", "bn", "AT1 and Tier 2 of BIL that count in the group (EUR bn)", "100%",
         "The ECB and the SRB. Whether the authority allows the derogation of Article 87(1) is open."),
        ("Fate of the senior debt of BIL", "Hold it (counts for 0% under SPE). Call it after an MREL disqualification event. Let it run off. Offer an exchange into debt of NBG.", "Calc N26", "=Calc!N26", "pct",
         "=Calc!N26*Calc!E24", "bn", "Senior debt of BIL that counts in the group (EUR bn)", "0%", "NBG and the holders. The terms of each instrument decide if a call is possible: not verified."),
        ("Internal MREL of BIL under SPE", "NBG buys instruments that BIL issues, ranked below the other liabilities of BIL. There is no scaling and no waiver between Greece and Luxembourg.", f"{MREL} section D",
         f"={q(MREL)}!B{ml['s_int_gap']}", "bn", f"={q(MREL)}!B{ml['s_int_gap']}", "bn", "Instruments that NBG must buy in BIL (EUR bn)", "As the model computes it",
         "The SRB sets it in the MREL decision of the resolution group."),
        ("Add-on of the NBG requirement for the stake in BIL under MPE", "0% to 100% of the price paid for the stake.", f"{MREL} section D", f"={q(MREL)}!C{ml['m_add_pct']}", "pct",
         f"={q(MREL)}!C{ml['m_add']}", "bn", "Add-on of the requirement (EUR bn)", "100%", "The SRB. The policy says that exposures between resolution groups raise the requirement (paragraphs 41 and 42)."),
        ("Valuation basis of BIL", "1 = tangible equity, 2 = CET1, 3 = average of the two.", "Calc N18", "=Calc!N18", "int", "=Calc!E167", "bn", "Value of BIL used (EUR bn)", "3",
         "NBG. CET1 is below tangible equity by the other CET1 deductions."),
        ("Number of nearest peers (k)", "From 10 to 40. A smaller k follows the return of BIL more closely and has more noise.", "Calc N17", "=Calc!N17", "int", "=Calc!E163", "x", "Median P/TBV of the nearest peers (x)", "20",
         "NBG. Section 7 of the sheet Valuation shows the effect of k."),
        ("Cost of equity of BIL", "0% = the cost of equity of the peers. 10% to 11% reflects a smaller, unlisted bank.", "Calc N24", "=Calc!N24", "pct", "=Calc!E231", "bn",
         "Effect on the peer value (EUR bn)", "10%", "NBG. The peers do not show a size premium: this is a judgment."),
        ("Growth rate g of the Gordon model", "0% to 3%.", "Calc N23", "=Calc!N23", "pct", "=Calc!E228", "f3", "Factor on the value on tangible equity", "2%", "NBG. Section 8 of the sheet Valuation shows the effect of g."),
        ("Peer screen", "Keep the rules (no subsidiaries, cooperatives, specialists, scale errors). Add a removed bank. Remove a bank by hand.", "Sheet Peers", f"=COUNT(Peers!D{W.PEER_TOP}:D{W.PEER_TOP + max(len(market['peers']['rows']), 1) - 1})", "int",
         "=Calc!E167", "bn", "Value of BIL used (EUR bn)", "Keep the rules", "NBG. Section 2 of the sheet Valuation shows the funnel."),
        ("Forward profit of BIL", "1 = from the drivers. 0 = flat at the trailing profit. BIL has no consensus.", f"{PNL_SHEET} B4", f"={q(PNL_SHEET)}!B4", "int", by_pnl("np"), "m1",
         "Profit of BIL in the analysis year: drivers less flat (EUR m)", "1", "NBG. The drivers are judgments: change the blue cells on the sheet P&L."),
        ("Synergies on the costs of BIL", "0% to 40% of the cost base of BIL, phased in over the years.", "Merger deal 1", f"={mg('syn_cost', 'D')}", "pct",
         by_year(lambda c: mg("e_imp", c)), "pct", "EPS accretion of NBG in the analysis year", "30%", "NBG management. BIL is cross-border: cost synergies are harder to reach than in an in-market deal."),
        ("Integration cost", "Share of the cost base of BIL, spread over the years.", "Calc N4", "=Calc!N4", "pct", "=Calc!H11", "bn", "Cost in the analysis year, after tax (EUR bn)", "20%", "NBG management."),
        ("Revenue attrition after the deal", "0% to 5% of the revenue of BIL.", "Merger deal 1", f"={mg('syn_rev', 'D')}", "pct", by_year(lambda c: mg("a_rv", c)), "m1",
         "Revenue effect after tax in the analysis year (EUR m)", "-3%", "NBG management. Customers can leave after a change of owner."),
        ("CET1 target and MREL target of NBG", "13% and 29% of TREA in the model. The MREL requirement of the group is lower.", "Calc D1 and D2", "=Calc!D1", "pct", "=Calc!H195", "bn",
         "Issuance need at the targets (EUR bn)", "13% and 29%", "NBG board and the ALCO. The need at the requirements is on the sheet MREL scenario."),
        ("Subordination requirement of the group", "13.5% of TREA is the most frequent filed value for a bank of EUR 100bn to 250bn. The SRB can set another.", "Calc N20", "=Calc!N20", "pct", "=Calc!H182", "bn",
         "Subordination requirement (EUR bn)", "13.5%", "The SRB. The date of the EUR 100bn test is not in the legal text: not verified."),
        ("Date of the EUR 100bn asset test", "At closing, or at the next resolution planning cycle. A transition of 3 years follows a new crossing.", "none: Calc N21", "=Calc!N21", "int",
         "=Calc!H7", "bn", "Assets of the group against the limit in N21 (EUR bn)", "The rule applies at closing", "The SRB. Not verified."),
    ]
    fmt = {"pct": "0.0%", "int": "0", "bn": "0.00", "x": '0.00"x"', "m1": "#,##0.0", "f3": "0.000"}
    for i, (dec, options, lever, now, nfmt, eff, efmt, elabel, default, who) in enumerate(items):
        r = 5 + i
        for j, v in enumerate([i + 1, dec, options, lever, now, eff, elabel, default, who, "Open"]):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=5).number_format = fmt[nfmt]
        ws.cell(row=r, column=5).font = GREEN
        ws.cell(row=r, column=6).number_format = fmt[efmt]
        ws.cell(row=r, column=6).font = GREEN
    ws.freeze_panes = "C5"


# ================================================================================================ Sources
def write_sources(ws, banks: list, deals: list, market: dict, pack: dict) -> None:
    name = deals[0]["short"]
    for i, w in enumerate((44, 56, 34, 100)):
        ws.column_dimensions[chr(65 + i)].width = w
    ws["A1"] = "Sources: where each number comes from, its date, and its limits."
    ws["A1"].font = Font(bold=True)
    pulled = market.get("pulled") or {}
    stamp = lambda k: str(pulled.get(k, "n/a"))[:10]  # noqa: E731
    te_n, te_t = sorted(pack.get("te_nbg") or {}), sorted(pack.get("te_tgt") or {})
    ecb = pack.get("ecb") or {}
    ecb_last = max((v["period"] for v in ecb.values() if isinstance(v, dict) and "period" in v), default="n/a")
    cc2n, cc2t = (pack.get("cc2_nbg") or {}).get("date", "n/a"), (pack.get("cc2_tgt") or {}).get("date", "n/a")
    or2n, or2t = sorted(pack.get("or2_nbg") or {}), sorted(pack.get("or2_tgt") or {})
    table = [
        ("Capital, buffers, leverage, liquidity (KM1)", "P3DH clean layer: clean_period_fact, template K_61.00", f"As of {banks[0].as_of} (NBG), {banks[1].as_of} ({name})", "Quarterly filings of the banks. The model uses the newest filing of each bank."),
        ("MREL resources and requirements (KM2)", "P3DH clean layer: template K_90.01", f"As of {banks[0].as_of} (NBG), {banks[1].as_of} ({name})", "KM2 gives own funds and eligible liabilities by resolution group. It does not give the issuer of each instrument."),
        ("Own funds detail, intangible assets (CC1), balance sheet lines (CC2)", "P3DH clean_fact: templates K_66.01 and K_66.02", f"CC2 of the lines: NBG {cc2n}, {name} {cc2t}", "CC2 gives the balance sheet in the accounting scope. The lines come from keywords and must tie to total assets within 2%."),
        ("Risk-weighted assets by type (OV1)", "P3DH clean_fact: template K_60.00", f"As of {banks[0].as_of}", "OV1 can have lines that the sheet leaves out."),
        ("Interest rate risk in the banking book (IRRBB1)", "P3DH clean_fact: template K_68.00", f"As of {banks[0].as_of}", "The sheet P&L uses the shift of 200 bp for net interest income."),
        ("Loans, non-performing loans, coverage (CR1)", "P3DH clean_fact: K_21.01 sheets a and d", "Newest filing of each bank", "CR1 sheet a gives the loans. The sheet d gives the impairment. Some filings leave sheet d empty."),
        ("Business indicator lines (OR2)", "P3DH clean_fact: K_19.02", f"Years: NBG {or2n[0] if or2n else 'n/a'} to {or2n[-1] if or2n else 'n/a'}; {name} {or2t[0] if or2t else 'n/a'} to {or2t[-1] if or2t else 'n/a'}",
         "OR2 has income lines only. It has no costs, impairment, or profit."),
        ("Costs, impairment, profit, history", "EBA Transparency Exercise TE2025: te_oth (public EBA data)", f"Reference dates: NBG {', '.join(te_n) or 'n/a'}; {name} {', '.join(te_t) or 'n/a'}",
         "P&L amounts are year to date. The last data is 30 June 2025. The EBA has stopped the exercise. Treat the data as history only."),
        ("Price, shares, market cap, trailing profit, opex, tax rate", "Bloomberg tables of p3dh.sqlite: bbg_entity, bbg_equity_snapshot, bbg_equity_daily", f"Pulled: equity {stamp('equity')}, entities {stamp('entities')}",
         "The per-share fields are in EUR and the price is in the currency of the stock. The model corrects for this (sheet Valuation, definitions)."),
        ("Consensus of NBG: net income, EPS, DPS, revenue, operating profit, ROE, BPS", "Bloomberg table bbg_consensus (BEST_ fields)", f"Pulled: equity {stamp('equity')}",
         f"{name} has no analyst consensus. Revenue and operating profit of NBG have fewer analysts than net income."),
        ("Deposit rate path and OIS curve", "Bloomberg tables bbg_rate_daily, bbg_expected_dfr", f"Rates pulled {stamp('rates')}", "Realised average to the pull date, then the OIS forward rate plus a spread."),
        ("Peer multiples", "Bloomberg snapshot (per-share, price to book) and P3DH KM1 (CET1)", f"Pulled: equity {stamp('equity')}", "Peer funnel on the sheet Valuation. The screen removes subsidiaries, cooperatives, specialists, and scale errors."),
        ("Luxembourg drivers (loan growth, new-loan spread, deposit beta, lending survey)", "ECB data (MIR, BSI, BLS) in ext_* tables", f"Latest month: {ecb_last}", "A country statistic, not data of BIL. The sheet P&L shows it as a note beside the judgment."),
        (f"Bonds of {name} by rank", "Bloomberg table bbg_bond", f"Pulled: {stamp('bonds')}", "Only EUR bonds are summed. A sum that mixes currencies is wrong. The senior stack in the model comes from KM2."),
        ("Regulatory rules (CRR, BRRD, SRB policy)", "Texts of the law and the SRB MREL policy; checked against the official texts", "Legal checks of 2026", "The sheet MREL scenario lists each statement with its status: Verified, Inference, or Unverified."),
    ]
    heads = ["Item", "Table or field", "Date or period", "Limit and quality"]
    for j, h in enumerate(heads):
        c = ws.cell(row=3, column=1 + j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD
    for i, row in enumerate(table):
        for j, v in enumerate(row):
            c = ws.cell(row=4 + i, column=1 + j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 4 + len(table) + 1
    ws.cell(row=r, column=1, value="Data caveats from the audit of the database").font = Font(bold=True)
    for k, line in enumerate((
        "Scale errors exist in P3DH or Bloomberg for a few banks. The peer screen removes a bank when its TREA differs from the RWA of Bloomberg by more than a factor of 2.",
        "Some peers were wrong when the model used the price in the currency of the stock. The model now uses the price to book ratio.",
        "The interest income sensitivity (IRRBB) had a sign error for banks with positive expense in OR2. It is corrected.",
        "A bank that files in thousands, or that reports a different scope, can look like an outlier. The sheet Peers lists every peer.",
        f"The CC2 lines of {name} exist for 31 December 2025 only. The EBA P&L of {name} ends on 30 June 2025.",
        "Bloomberg bond amounts mix currencies. The model does not use them for MREL.",
        f"The Bloomberg consensus covers NBG only. The forward profit of {name} is a judgment (sheet P&L).",
        "The workbook has no stored values. Excel calculates on opening. The Streamlit page calculates with the same formulas.",
    )):
        ws.cell(row=r + 1 + k, column=1, value=line)
    r2 = r + 2 + 8
    if pack.get("missing"):
        ws.cell(row=r2, column=1, value="Data that the build could not read").font = Font(bold=True)
        for k, m in enumerate(pack["missing"]):
            ws.cell(row=r2 + 1 + k, column=1, value=m)
    ws.freeze_panes = "A4"


# ================================================================================================ all sheets
def write_sheets(wb, banks: list, deals: list, g: dict, market: dict) -> None:
    pack = market.get("pack") or {}
    sheets = {n: wb.create_sheet(n) for n in NEW_SHEETS}
    rows: dict[str, dict[str, int]] = {}
    rows[PNL_SHEET] = ma_pack_pnl.write_pnl(sheets[PNL_SHEET], banks, market, pack)
    rows[MREL] = ma_pack_pnl.write_mrel(sheets[MREL], banks, market, pack)
    rows[VAL] = write_valuation(sheets[VAL], banks, deals, g, market, pack, rows[PNL_SHEET])
    rows[BS] = write_balance(sheets[BS], banks, deals, market, pack)
    write_decisions(sheets[DEC], banks, deals, market, pack, rows)
    write_sources(sheets[SRC], banks, deals, market, pack)
