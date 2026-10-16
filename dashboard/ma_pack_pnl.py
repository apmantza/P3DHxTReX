"""Sheets P&L and MREL scenario of the deal pack. See ma_pack.py for the other sheets. Labels and notes follow ASD-STE100."""

from __future__ import annotations

from typing import Any

import ma_workbook as W
from openpyxl.styles import Font

from ma_pack_base import BIL_KEYS, BIL_TOP, BLUE, INPUT_FILL, PNL_SWITCH, PNL_YC, Sheet, bil_row

NOTE_FLAT = "Flat case: the trailing profit from continuing operations (Bloomberg) for every year, as the model had before."


def _lines(d: dict, year: int) -> dict[str, float]:
    return d.get(year, {})


def write_pnl(ws, banks: list, market: dict, pack: dict) -> dict[str, int]:
    S = Sheet(ws, (70, 14, 14, 14, 14, 14, 14, 100))
    k_ = W.MCOLS[1]  # column of BIL on sheet Market
    ws["A1"] = "P&L of NBG and BIL: history, forwards, and the combined profit. EUR million. Green: data. Blue on yellow: input. Black: formula."
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "BIL has no analyst consensus. Its forward profit comes from drivers (switch below) or stays flat at the trailing profit. NBG forwards come from the Bloomberg consensus."
    ws["A4"] = "BIL forward profit: 1 = from the drivers, 0 = flat at the trailing profit (Bloomberg, continuing operations)"
    ws[PNL_SWITCH] = 1
    ws[PNL_SWITCH].font, ws[PNL_SWITCH].fill = BLUE, INPUT_FILL
    ws["A5"] = "The switch feeds the revenue, costs, and profit of BIL on the sheet Merger (deal 1)."
    S.r = 6
    S.head(["BIL: forward profit (EUR m)", "2025A (base)"] + W.YEARS + ["Source and note"])
    S.section("Inputs of each year (judgments: change them to test)")
    S.r = BIL_TOP
    or2, te = pack.get("or2_tgt", {}), pack.get("te_tgt", {})
    years = sorted(or2)
    y25 = or2.get(2025) or (or2[years[-1]] if years else {})
    nii25 = (y25.get("interest_income", 0) - y25.get("interest_expense", 0)) if y25 else None
    fees25 = (y25.get("fee_income", 0) - y25.get("fee_expense", 0)) if y25 else None
    other25 = (y25.get("dividends", 0) + y25.get("other_income", 0) - y25.get("other_expense", 0) + y25.get("trading", 0) + y25.get("banking_book", 0)) if y25 else None
    cr1 = pack.get("cr1_tgt", {})
    hist = cr1.get("history", {}) if cr1 else {}
    base_date = "2025-12-31" if "2025-12-31" in hist else (cr1.get("date") if cr1 else None)
    loans0 = (hist.get("2025-12-31") or cr1.get("gross")) if cr1 else None
    imp24 = te.get("2024-12-31", {}).get("lines", {}).get("impairment")
    cor_default = round(imp24 / loans0 * 10000) if imp24 and loans0 else 20
    h1 = te.get("2025-06-30", {})
    h1_bp = (h1["lines"]["impairment"] * 4 / h1["n_quarters"] / loans0 * 10000) if loans0 and h1.get("lines", {}).get("impairment") is not None and h1.get("n_quarters") else None
    ecb = pack.get("ecb", {})
    slot = market["slots"][1]
    d1 = max(hist) if hist else None
    d0 = f"{int(d1[:4]) - 1}{d1[4:]}" if d1 else None
    grew = f" CR1: the loans grew {hist[d1] / hist[d0] - 1:.1%} in the 12 months to {d1}." if d1 and d0 in hist else ""

    def ecb_note(label: str, short: str) -> str:
        v = ecb.get(label)
        if not v:
            return ""
        eu = f" (euro area {v['euro_area']:.1f}%)" if v.get("euro_area") is not None else ""
        return f" ECB, Luxembourg, {short}, {v['period']}: {v['value']:.1f}%{eu}."

    spread = ecb.get("new corporate loans: rate less 3-month Euribor (points)")
    if spread and spread.get("before") is not None:
        change_bp = (spread["value"] - spread["before"]) * 100
        m1 = round(0.2 * change_bp)
        margin = [m1, m1, round(m1 / 2), round(m1 / 2)]
        margin_note = (f"Assumption: a fifth of the book reprices each year. Year 1 and 2: a fifth of the change in the spread of new corporate loans in Luxembourg over 12 months "
                       f"({spread['before']:.2f} to {spread['value']:.2f} points above the 3-month Euribor). Year 3 and 4: half of that.")
    else:
        margin, margin_note = [-5] * 4, "Judgment. The ECB spread of new corporate loans is not loaded."
    tax_eff, tax_stat = slot.get("tax_eff"), slot.get("tax_stat")
    tax_note = ("Effective rate of Bloomberg, else the statutory rate. " + (f"Effective {tax_eff:.1%}, statutory {tax_stat:.1%}: the statutory rate changes the profit by {(1 - tax_stat) / (1 - tax_eff) - 1:+.0%}. " if tax_eff and tax_stat else "")
                + "Overwrite to test.")
    cost_te = [te.get("2024-12-31", {}).get("lines", {}).get(k) for k in ("admin", "depreciation", "resolution")]
    cost_note = ("Base: operating expenses of the last fiscal year (Bloomberg)." + (f" The EBA exercise for 2024 gives administrative expenses, depreciation, and contributions of {sum(abs(x) for x in cost_te):,.0f}." if None not in cost_te else ""))

    y4 = lambda v: [None] + list(v)  # noqa: E731
    S.row("in_loan_g", "Growth of the loan stock", y4([0.04, 0.03, 0.03, 0.03]), "pct", "Judgment." + grew + ecb_note("loan growth, corporations (%)", "loans to corporations") + ecb_note("loan growth, house purchase (%)", "house purchase loans"), ["input"] * 5)
    S.row("in_margin_bp", "Change of the margin on the average loans (bp a year)", y4(margin), "bp", margin_note, ["input"] * 5)
    S.row("in_fee_g", "Growth of net fees", y4([0.03] * 4), "pct", "Judgment.", ["input"] * 5)
    S.row("in_cost_g", "Growth of costs", y4([0.03] * 4), "pct", "Judgment. The sheet Merger uses 3% for the cost base of the targets.", ["input"] * 5)
    S.row("in_cor_bp", "Cost of risk (impairment, bp of the average loans)", y4([cor_default] * 4), "bp",
          ("Default: impairment of the financial assets in the EBA exercise for 2024 / loans." + (f" For the first half of 2025, annualised, it is {h1_bp:.0f} bp: the cost of risk varies a lot." if h1_bp is not None else "")) if imp24 else "Default 20 bp: the EBA exercise is not loaded.", ["input"] * 5)
    tax_f = f"=IF(Market!{k_}{W.MROW['tax_eff']}>0,Market!{k_}{W.MROW['tax_eff']},Market!{k_}{W.MROW['tax']})"
    S.row("in_tax", "Tax rate", y4([tax_f] * 4), "pct", tax_note, None)
    S.row("in_other", "Other income: dividends, other operating, trading, banking book (flat)", y4([round(other25, 1) if other25 is not None else 0.0] * 4), "m1", "Default: the value of 2025 (OR2). Held flat.", [None] + ["data"] * 4)
    S.row("gap1", "Drivers", [], None, "", None, bold=True)
    S.row("loans", "Gross loans, end of year (CR1)", [loans0 if loans0 is not None else 0.0] + [f"={p}{bil_row('loans')}*(1+{c}{bil_row('in_loan_g')})" for p, c in zip(["B"] + PNL_YC[:-1], PNL_YC)], "m", f"Base: gross loans at {base_date or 'n/a'} (CR1).", ["data"])
    S.row("loans_avg", "Average gross loans", [None] + [f"=({p}{bil_row('loans')}+{c}{bil_row('loans')})/2" for p, c in zip(["B"] + PNL_YC[:-1], PNL_YC)], "m")
    S.row("dfr", "Average deposit rate of the year (ECB deposit facility rate: realised and OIS forward)", [pack.get("dfr_2025")] + [f"=Market!{m}{W.MKT_DFR_ROW}" for m in W.MKT_YC], "pct2", "Sheet Market, row 87.", ["data"])
    S.row("d_dfr", "Change of the average rate against the year before", [None] + [f"={c}{bil_row('dfr')}-{p}{bil_row('dfr')}" for p, c in zip(["B"] + PNL_YC[:-1], PNL_YC)], "pct2")
    up, dn = pack.get("irrbb_nii_tgt", {}).get("0010"), pack.get("irrbb_nii_tgt", {}).get("0020")
    S.row("irrbb_up", "IRRBB1: change of net interest income for a parallel shift of +200 bp (EUR m)", [up[0] / 1e6 if up else 0.0], "m1", f"P3DH K_68.00, {up[1] if up else 'no data'}.", ["data"])
    S.row("irrbb_dn", "IRRBB1: change of net interest income for a parallel shift of -200 bp (EUR m)", [dn[0] / 1e6 if dn else 0.0], "m1", f"P3DH K_68.00, {dn[1] if dn else 'no data'}.", ["data"])
    S.row("gap2", "Net interest income and profit", [], None, "", None, bold=True)
    prev = ["B"] + PNL_YC[:-1]
    S.row("nii_vol", "Net interest income: volume effect (previous net interest income x loan growth)", [None] + [f"={p}{bil_row('nii')}*{c}{bil_row('in_loan_g')}" for p, c in zip(prev, PNL_YC)], "m1")
    S.row("nii_rate", "Net interest income: rate effect (IRRBB1 sensitivity x change of the average rate / 2 points)",
          [None] + [f"=IF({c}{bil_row('d_dfr')}>=0,$B${bil_row('irrbb_up')}*{c}{bil_row('d_dfr')}/0.02,$B${bil_row('irrbb_dn')}*(-{c}{bil_row('d_dfr')})/0.02)" for c in PNL_YC], "m1",
          "A rise of the rate uses the +200 bp figure. A fall uses the -200 bp figure. The figure is linear in the size of the move.")
    S.row("nii_margin", "Net interest income: margin effect (margin change x average loans)", [None] + [f"={c}{bil_row('in_margin_bp')}/10000*{c}{bil_row('loans_avg')}" for c in PNL_YC], "m1")
    S.row("nii", "Net interest income", [nii25 if nii25 is not None else 0.0] + [f"={p}{bil_row('nii')}+{c}{bil_row('nii_vol')}+{c}{bil_row('nii_rate')}+{c}{bil_row('nii_margin')}" for p, c in zip(prev, PNL_YC)], "m1",
          "Base: interest income less interest expense of 2025 (OR2).", ["data"], bold=True)
    S.row("fees", "Net fees", [fees25 if fees25 is not None else 0.0] + [f"={p}{bil_row('fees')}*(1+{c}{bil_row('in_fee_g')})" for p, c in zip(prev, PNL_YC)], "m1", "Base: fee income less fee expense of 2025 (OR2).", ["data"])
    S.row("other", "Other income", [other25 if other25 is not None else 0.0] + [f"={c}{bil_row('in_other')}" for c in PNL_YC], "m1", "", ["data"])
    S.row("rev", "Total revenue", [f"={c}{bil_row('nii')}+{c}{bil_row('fees')}+{c}{bil_row('other')}" for c in ["B"] + PNL_YC], "m1", "", None, bold=True)
    S.row("costs", "Costs (negative)", [f"=-Market!{k_}{W.MROW['opex']}"] + [f"={p}{bil_row('costs')}*(1+{c}{bil_row('in_cost_g')})" for p, c in zip(prev, PNL_YC)], "m1", cost_note)
    S.row("impair", "Impairment (negative)", [-imp24 if imp24 else 0.0] + [f"=-{c}{bil_row('in_cor_bp')}/10000*{c}{bil_row('loans_avg')}" for c in PNL_YC], "m1", "Base: impairment of the financial assets in 2024 (EBA exercise).", ["data"])
    S.row("pbt", "Profit before tax", [f"={c}{bil_row('rev')}+{c}{bil_row('costs')}+{c}{bil_row('impair')}" for c in ["B"] + PNL_YC], "m1")
    S.row("tax", "Tax", [f"=-B{bil_row('pbt')}*C{bil_row('in_tax')}"] + [f"=-{c}{bil_row('pbt')}*{c}{bil_row('in_tax')}" for c in PNL_YC], "m1")
    S.row("np_drv", "Net profit from the drivers", [f"={c}{bil_row('pbt')}+{c}{bil_row('tax')}" for c in ["B"] + PNL_YC], "m1", "", None, bold=True)
    S.row("np_flat", "Net profit flat (trailing, continuing operations)", [f"=Market!{k_}{W.MROW['ni_base']}"] + [f"=Market!{k_}{W.MROW['ni_base']}" for _ in PNL_YC], "m1", NOTE_FLAT)
    S.row("np_sel", "Net profit used by the model (switch in B4)", [f"=B{bil_row('np_flat')}"] + [f"=IF({PNL_SWITCH.replace('B', '$B$')}=1,{c}{bil_row('np_drv')},{c}{bil_row('np_flat')})" for c in PNL_YC], "m1", "", None, bold=True)
    S.row("gap3", "", [])
    S.row("check_np", "Check of the base year: profit from the lines less the trailing continuing profit (EUR m)", [f"=B{bil_row('np_drv')}-B{bil_row('np_flat')}"], "m1",
          "A small gap means the lines explain the reported profit. A large gap means a line is missing or the tax rate is wrong.")
    assert S.rows == {k: BIL_TOP + i for i, k in enumerate(BIL_KEYS)} or all(S.rows[k] == bil_row(k) for k in BIL_KEYS), "layout of the BIL block differs from BIL_KEYS"
    S.skip()
    S.row("rote_fwd", "RoTE of BIL on the tangible equity of today (profit used / tangible equity)", [f"=B{bil_row('np_sel')}/1000/Calc!$E$30"] + [f"={c}{bil_row('np_sel')}/1000/Calc!$E$30" for c in PNL_YC], "pct",
          "The tangible equity of today is held flat: a simple measure. Retained profit would lower the return.")
    S.skip()

    # ---------------------------------------------------------------- combined P&L (sheet Merger, deal 1)
    R = W.merger_start(0)
    mo = lambda key, col: f"Merger!{col}{R + W.MOFF[key]}"  # noqa: E731
    S.head(["Combined P&L of NBG and BIL (EUR m), from the sheet Merger, deal 1", "", *W.YEARS, "", "Source and note"])
    items = [("b_rev", "NBG revenue (consensus)", "m"), ("t_rev", "BIL revenue", "m"), ("b_cost", "NBG costs", "m"), ("t_cost", "BIL costs", "m"), ("b_np", "NBG net profit (consensus)", "m"),
             ("t_np", "BIL net profit (100%)", "m"), ("a_tnp", "BIL net profit (stake)", "m"), ("a_cs", "Cost synergies after tax", "m"), ("a_rs", "Restructuring after tax", "m"),
             ("a_rv", "Revenue synergies or attrition after tax", "m"), ("a_fin", "Deal financing after tax (cash paid at the deposit rate)", "m"), ("a_tot", "Pro forma net profit", "m"),
             ("s_tot", "Shares after the deal (million)", "m1"), ("e_eps0", "EPS of NBG before the deal (EUR)", "eps"), ("e_eps1", "EPS pro forma (EUR)", "eps"), ("e_imp", "EPS accretion or dilution", "pct"),
             ("v_rote", "RoTE pro forma", "pct"), ("e_roi", "Return on investment", "pct"), ("c_ratio", "CET1 ratio pro forma", "pct"), ("c_head", "CET1 headroom to the MDA trigger (bp)", "bp")]
    for key, label, fmt in items:
        S.row("c_" + key, label, [None] + [f"={mo(key, c)}" for c in W.YC], fmt, "", None, bold=key in ("a_tot", "e_eps1", "e_imp"))
    S.row("c_cir", "Cost/income ratio of the group before synergies", [None] + [f"=-({mo('b_cost', c)}+{mo('t_cost', c)})/({mo('b_rev', c)}+{mo('t_rev', c)})" for c in W.YC], "pct")
    S.row("c_cir_syn", "Cost/income ratio of the group after cost synergies and attrition", [None] + [f"=-({mo('b_cost', c)}+{mo('t_cost', c)}+{mo('t_cost_syn', c)})/({mo('b_rev', c)}+{mo('t_rev', c)}+{mo('t_rev_syn', c)})" for c in W.YC], "pct")
    S.skip()

    # ---------------------------------------------------------------- NBG forwards
    S.head(["NBG: Bloomberg consensus (EUR m unless stated)", "", *W.YEARS, "", "Source and note"])
    cons = market["slots"][0].get("cons", {})
    mc = lambda m, k: f"Market!B{W.cons_row(m, k)}"  # noqa: E731
    S.row("n_ni", "Net income", [None] + [f"={mc('NI', k)}" for k in range(4)], "m", "Sheet Market (BEST_NET_INCOME).")
    S.row("n_eps", "EPS (EUR)", [None] + [f"={mc('EPS', k)}" for k in range(4)], "eps")
    S.row("n_dps", "DPS (EUR)", [None] + [f"={mc('DPS', k)}" for k in range(4)], "eps")
    S.row("n_sales", "Revenue", [None] + [f"={mc('SALES', k)}" for k in range(4)], "m")
    S.row("n_opp", "Operating profit", [None] + [f"={mc('OPP', k)}" for k in range(4)], "m", "No operating profit consensus for NBG: the costs of NBG are opex of the last fiscal year grown at 3% on the sheet Merger.")
    S.row("n_cost", "Costs implied (revenue less operating profit)", [None] + [f"=IF(AND({c}{S.rows['n_sales']}<>0,{c}{S.rows['n_opp']}<>0),-({c}{S.rows['n_sales']}-{c}{S.rows['n_opp']}),0)" for c in PNL_YC], "m")
    S.row("n_roe", "Return on equity (BEST_ROE, %)", [None] + [cons.get(2026 + k, {}).get("ROE") for k in range(4)], "m1", "Data of Bloomberg.", [None] + ["data"] * 4)
    S.row("n_bps", "Book value per share (BEST_BPS, EUR)", [None] + [cons.get(2026 + k, {}).get("BPS") for k in range(4)], "eps", "Data of Bloomberg.", [None] + ["data"] * 4)
    S.row("n_pay", "Payout: DPS / EPS", [None] + [f"=IF({c}{S.rows['n_eps']}<>0,{c}{S.rows['n_dps']}/{c}{S.rows['n_eps']},0)" for c in PNL_YC], "pct")
    S.row("n_sh", "Shares implied (million): net income / EPS", [None] + [f"=IF({c}{S.rows['n_eps']}<>0,{c}{S.rows['n_ni']}/{c}{S.rows['n_eps']},0)" for c in PNL_YC], "m1")
    S.row("n_g", "Growth of net income", [None, None] + [f"={c}{S.rows['n_ni']}/{p}{S.rows['n_ni']}-1" for p, c in zip(PNL_YC[:-1], PNL_YC[1:])], "pct")
    S.skip()

    # ---------------------------------------------------------------- history, OR2 (P3DH)
    for nm, key, bank in (("NBG", "or2_nbg", banks[0]), ("BIL", "or2_tgt", banks[1])):
        d = pack.get(key, {})
        ys = sorted(d)[-3:]
        S.head([f"{nm}: business indicator lines of OR2 (P3DH, EUR m)", *[f"FY{y}" for y in ys], "", "", "", "Source and note"] if ys else [f"{nm}: OR2 not available", "", "", "", "", "", "", ""])
        if not ys:
            S.skip()
            continue
        col = lambda i: chr(66 + i)  # noqa: E731
        g = lambda name, sign=1: [sign * d[y].get(name, 0.0) for y in ys]  # noqa: E731
        p0 = f"{nm}_"
        S.row(p0 + "ii", "Interest and lease income", g("interest_income"), "m", "P3DH K_19.02 row 1a.", "data")
        S.row(p0 + "ie", "Interest and lease expense", g("interest_expense", -1), "m", "", "data")
        S.row(p0 + "nii", "Net interest income", [f"={col(i)}{S.rows[p0 + 'ii']}+{col(i)}{S.rows[p0 + 'ie']}" for i in range(len(ys))], "m", "", None, bold=True)
        S.row(p0 + "fi", "Fee and commission income", g("fee_income"), "m", "", "data")
        S.row(p0 + "fe", "Fee and commission expense", g("fee_expense", -1), "m", "", "data")
        S.row(p0 + "nf", "Net fees", [f"={col(i)}{S.rows[p0 + 'fi']}+{col(i)}{S.rows[p0 + 'fe']}" for i in range(len(ys))], "m", "", None, bold=True)
        S.row(p0 + "oo", "Other operating income, net, and dividends", [d[y].get("other_income", 0) - d[y].get("other_expense", 0) + d[y].get("dividends", 0) for y in ys], "m", "", "data")
        S.row(p0 + "tr", "Result of the trading book and the banking book", [d[y].get("trading", 0) + d[y].get("banking_book", 0) for y in ys], "m", "", "data")
        S.row(p0 + "inc", "Income before costs", [f"={col(i)}{S.rows[p0 + 'nii']}+{col(i)}{S.rows[p0 + 'nf']}+{col(i)}{S.rows[p0 + 'oo']}+{col(i)}{S.rows[p0 + 'tr']}" for i in range(len(ys))], "m", "P3DH has no costs, impairment, or profit.", None, bold=True)
        S.skip()

    # ---------------------------------------------------------------- history, EBA exercise
    S.text("EBA Transparency Exercise (public EBA data, TE2025: last data 30 June 2025; the EBA has stopped the exercise). P&L amounts are year to date: the row 'quarters' says how many quarters.", bold=True)
    for nm, key in (("NBG", "te_nbg"), ("BIL", "te_tgt")):
        d = pack.get(key, {})
        ps = sorted(d)
        if not ps:
            S.text(f"{nm}: not available. Run: python scripts/pull_ext.py eba")
            continue
        col = lambda i: chr(66 + i)  # noqa: E731
        S.head([f"{nm}: P&L (EBA exercise, EUR m)", *[p for p in ps], *([""] * (6 - len(ps))), "Source and note"])
        ln = lambda name, sign=1: [None if name not in d[p]["lines"] else sign * d[p]["lines"][name] for p in ps]  # noqa: E731
        p0 = f"te_{nm}_"
        S.row(p0 + "q", "Quarters of the fiscal year in the amount", [d[p]["n_quarters"] for p in ps], "int", "", "data")
        S.row(p0 + "ii", "Interest income", ln("interest_income"), "m1", "", "data")
        S.row(p0 + "ie", "Interest expenses", ln("interest_expense", -1), "m1", "", "data")
        S.row(p0 + "nii", "Net interest income", [f"={col(i)}{S.rows[p0 + 'ii']}+{col(i)}{S.rows[p0 + 'ie']}" for i in range(len(ps))], "m1", "", None, bold=True)
        S.row(p0 + "nf", "Net fees", ln("fees"), "m1", "", "data")
        S.row(p0 + "inc", "Total operating income, net", ln("income"), "m1", "", "data", bold=True)
        S.row(p0 + "adm", "Administrative expenses", ln("admin", -1), "m1", "", "data")
        S.row(p0 + "dep", "Depreciation", ln("depreciation", -1), "m1", "", "data")
        S.row(p0 + "res", "Contributions to resolution funds and deposit guarantee schemes", ln("resolution", -1), "m1", "", "data")
        S.row(p0 + "imp", "Impairment of financial assets", ln("impairment", -1), "m1", "", "data")
        S.row(p0 + "pbt", "Profit before tax from continuing operations", ln("pbt"), "m1", "", "data")
        S.row(p0 + "np", "Profit for the period", ln("profit"), "m1", "", "data", bold=True)
        S.row(p0 + "ass", "Total assets at the date", ln("assets"), "m", "", "data")
        S.row(p0 + "eq", "Total equity at the date", ln("equity"), "m", "", "data")
        S.row(p0 + "cir", "Cost/income ratio", [f"=IF({col(i)}{S.rows[p0 + 'inc']}>0,-({col(i)}{S.rows[p0 + 'adm']}+{col(i)}{S.rows[p0 + 'dep']}+{col(i)}{S.rows[p0 + 'res']})/{col(i)}{S.rows[p0 + 'inc']},0)" for i in range(len(ps))], "pct")
        S.row(p0 + "roe", "Return on equity, annualised (profit x 4 / quarters / equity)", [f"=IF({col(i)}{S.rows[p0 + 'eq']}>0,{col(i)}{S.rows[p0 + 'np']}*4/{col(i)}{S.rows[p0 + 'q']}/{col(i)}{S.rows[p0 + 'eq']},0)" for i in range(len(ps))], "pct")
        S.row(p0 + "cor", "Impairment / total assets, annualised (bp)", [f"=IF({col(i)}{S.rows[p0 + 'ass']}>0,-{col(i)}{S.rows[p0 + 'imp']}*4/{col(i)}{S.rows[p0 + 'q']}/{col(i)}{S.rows[p0 + 'ass']}*10000,0)" for i in range(len(ps))], "bp")
        S.skip()
    ws.freeze_panes = "B7"
    return S.rows


def write_mrel(ws, banks: list, market: dict, pack: dict) -> dict[str, int]:
    S = Sheet(ws, (74, 14, 14, 14, 14, 14, 14, 100))
    ws["A1"] = "MREL scenario of NBG + BIL. EUR billion unless stated. Green: data. Blue on yellow: input. Black: formula."
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "The SRB decides the resolution strategy alone. NBG has not been told. This sheet shows both strategies, with the strategy that you assume in B4."
    ws["A4"] = "Strategy assumed: 1 = single point of entry (NBG resolution entity, BIL a subsidiary), 2 = multiple points of entry (BIL stays its own resolution entity)"
    ws["B4"] = 1
    ws["B4"].font, ws["B4"].fill = BLUE, INPUT_FILL
    S.r = 6
    S.head(["A. Requirement", "NBG", "BIL", "NBG + BIL", "", "", "", "Source and note"])
    row = lambda key, label, vals, fmt, note="", roles=None, bold=False: S.row(key, label, vals, fmt, note, roles, bold)  # noqa: E731
    row("assets", "Total assets", ["=Calc!D7", "=Calc!E7", "=Calc!H7"], "bn", "P3DH. The combination includes the goodwill.")
    row("trea", "Risk-weighted assets (TREA)", ["=Calc!D8", "=Calc!E8", "=Calc!H8"], "bn")
    row("req_pct", "MREL requirement, % of TREA, without buffer (KM2 row EU-7; RWA-weighted for the combination)", ["=Calc!D119", "=Calc!E119", "=Calc!H119"], "pct2", "P3DH KM2. The user rule: the own funds requirement for MREL is the MREL requirement plus the combined buffer.")
    row("cbr", "Combined buffer requirement", ["=Calc!D99", "=Calc!E99", "=Calc!H99"], "pct2", "P3DH KM1.")
    row("req_tot", "MREL requirement with the combined buffer, % of TREA", ["=Calc!D120", "=Calc!E120", "=Calc!H176"], "pct2")
    row("req_bn", "MREL requirement with the combined buffer (EUR bn)", ["=Calc!D181", "=Calc!E181", "=Calc!H181"], "bn", "", None, True)
    row("sub_pct", "Subordination requirement with the combined buffer, % of TREA (filed, else the assumption N20 for a combination above N21)", ["=Calc!D177", "=Calc!E177", "=Calc!H177"], "pct2",
        "Top-tier rule (BRRD 45c(5)): resolution groups above EUR 100bn of assets need 13.5% of TREA and 5% of the leverage exposure in subordinated resources, with the buffer on top.")
    row("sub_bn", "Subordination requirement (EUR bn); the combination is not below N22 of the leverage exposure", ["=Calc!D182", "=Calc!E182", "=Calc!H182"], "bn", "", None, True)
    row("top", "Top-tier test: assets of the combination against the limit N21 (EUR bn)", [None, "=Calc!N21", "=Calc!H7"], "bn", "Column C holds the limit. Column D holds the assets of the combination (pro forma).")
    row("top_flag", "Result", [None, None, '=IF(Calc!H7>=Calc!N21,"Above the limit: top-tier rule","Below the limit")'], "text",
        "A group that newly crosses the limit has a transition of 3 years (BRRD 45m(4), SRMR 12k(4)). The date of the test is not in the text: unverified.")
    S.skip()
    S.head(["B. Resources", "NBG", "BIL", "NBG + BIL", "", "", "", "Source and note"])
    row("cet1", "CET1", ["=Calc!D20", "=Calc!E20", "=Calc!H20"], "bn")
    row("at1", "AT1 (the combination: AT1 of BIL at N19 x recognition of row 233)", ["=Calc!D22", "=Calc!E22", "=Calc!H22"], "bn")
    row("t2", "Tier 2 (the same recognition)", ["=Calc!D23", "=Calc!E23", "=Calc!H23"], "bn")
    row("own", "Own funds", [f"=B{S.rows['cet1']}+B{S.rows['at1']}+B{S.rows['t2']}", f"=C{S.rows['cet1']}+C{S.rows['at1']}+C{S.rows['t2']}", f"=D{S.rows['cet1']}+D{S.rows['at1']}+D{S.rows['t2']}"], "bn", "", None, True)
    row("sen", "Eligible liabilities (senior; the combination counts the senior of BIL at N26: 0% under single point of entry)", ["=Calc!D24", "=Calc!E24", "=Calc!H24"], "bn",
        "KM2 of P3DH: own funds and eligible liabilities less own funds. Bloomberg bond amounts mix currencies and are not used here.")
    row("res", "MREL resources", ["=Calc!D25", "=Calc!E25", "=Calc!H25"], "bn", "", None, True)
    row("ratio", "MREL ratio, % of TREA", ["=Calc!D26", "=Calc!E26", "=Calc!H26"], "pct")
    row("head", "Headroom to the requirement with the buffer (bp of TREA)", ["=Calc!D121", "=Calc!E121", "=Calc!H121"], "bp")
    row("surplus", "Surplus (+) or shortfall (-) against the requirement with the buffer (EUR bn)", [f"=B{S.rows['res']}-B{S.rows['req_bn']}", f"=C{S.rows['res']}-C{S.rows['req_bn']}", f"=D{S.rows['res']}-D{S.rows['req_bn']}"], "bn", "", None, True)
    row("subres", "Subordinated resources: own funds and senior non-preferred", ["=Calc!D186", "=Calc!E186", "=Calc!H186"], "bn")
    row("sub_head", "Surplus (+) or shortfall (-) against the subordination requirement (EUR bn)", [f"=IF(B{S.rows['sub_bn']}>0,B{S.rows['subres']}-B{S.rows['sub_bn']},0)", f"=IF(C{S.rows['sub_bn']}>0,C{S.rows['subres']}-C{S.rows['sub_bn']},0)", f"=IF(D{S.rows['sub_bn']}>0,D{S.rows['subres']}-D{S.rows['sub_bn']},0)"], "bn", "0 = no requirement.")
    S.skip()
    S.head(["C. Recognition of the AT1 and Tier 2 of BIL in the own funds of the group (CRR Article 87)", "BIL", "", "", "", "", "", "Source and note"])
    gross = W.p3dh_rows()["cc1_cet1_before_adj"]
    row("r_own", "Own funds of BIL", ["=Calc!E20+Calc!E22+Calc!E23"], "bn", "P3DH KM1.")
    row("r_req_own", "Own requirement of BIL: total capital requirement x TREA of BIL (8% + P2R + combined buffer)", ["=Calc!E102*Calc!E8"], "bn")
    row("r_req_grp", "Requirement that the group has for BIL: total capital requirement of NBG x TREA of BIL", ["=Calc!D102*Calc!E8"], "bn", "The authority may allow the higher of the two amounts instead of the lower (Article 87(1), last subparagraph).")
    row("r_low", "Lower of the two requirements", [f"=MIN(B{S.rows['r_req_own']},B{S.rows['r_req_grp']})"], "bn")
    row("r_sur", "Surplus of BIL over that requirement", [f"=MAX(0,B{S.rows['r_own']}-B{S.rows['r_low']})"], "bn", "The share of the surplus that belongs to outside investors in BIL cannot absorb losses elsewhere in the group.")
    row("r_items", "Gross capital items of BIL: CET1 before adjustments (CC1) + AT1 + Tier 2", [f"=P3DH!E{gross}+Calc!E22+Calc!E23"], "bn")
    row("r_factor", "Share recognised = 1 - surplus / gross capital items", [f"=IF(B{S.rows['r_items']}<=0,1,1-B{S.rows['r_sur']}/B{S.rows['r_items']})"], "pct", "The same as Calc row 233. The factor does not depend on the stake: it applies to every outside-held own funds item of BIL (minority CET1, AT1, Tier 2).", None, True)
    row("r_chk", "Check: equal to Calc row 233 (1 = yes)", [f"=IF(ABS(B{S.rows['r_factor']}-Calc!E233)<0.000001,1,0)"], "flag")
    row("r_out", "AT1 and Tier 2 of BIL outstanding", ["=Calc!E22+Calc!E23"], "bn")
    row("r_rec", "Recognised in the own funds of the group (x N19)", [f"=B{S.rows['r_out']}*Calc!E233*Calc!N19"], "bn")
    row("r_exc", "Not recognised", [f"=B{S.rows['r_out']}-B{S.rows['r_rec']}"], "bn", "", None, True)
    S.skip()
    S.head(["D. Result by strategy", "Single point of entry", "Multiple points of entry", "Selected (B4)", "", "", "", "Source and note"])
    S.section("Single point of entry: BIL is a subsidiary in the resolution group of NBG")
    row("s_req", "Requirement of the group with the buffer (EUR bn)", ["=Calc!H181"], "bn")
    row("s_res", "Resources of the group (EUR bn)", ["=Calc!H25"], "bn", "BIL senior counts at N26 (0%).")
    row("s_sur", "Surplus (+) or shortfall (-) against the requirement (EUR bn)", [f"=B{S.rows['s_res']}-B{S.rows['s_req']}"], "bn", "", None, True)
    row("s_need_req", "Issuance need at the requirements (Calc row 192, EUR bn)", ["=Calc!H192"], "bn")
    row("s_need_tgt", "Issuance need at the targets D1 and D2 (Calc row 195, EUR bn)", ["=Calc!H195"], "bn", "The larger of the need at the targets and at the requirements.", None, True)
    row("s_strand", "Senior and senior non-preferred debt of BIL that counts nowhere (EUR bn)", ["=Calc!E24"], "bn",
        "It does not count for the group: only liabilities that NBG issues count (SRB MREL policy, paragraph 111). It does not count for BIL internal MREL: that needs instruments that NBG buys (BRRD 45f(2)).")
    row("s_int_req", "Internal MREL of BIL: requirement with the buffer (EUR bn)", ["=Calc!E181"], "bn", "No scaling of the requirement and no waiver between Greece and Luxembourg (SRB MREL policy, paragraph 37; BRRD 45f(3)).")
    row("s_int_res", "Internal MREL of BIL: its own funds (EUR bn)", ["=Calc!E20+Calc!E22+Calc!E23"], "bn", "Outside-held AT1 and Tier 2 count for internal MREL if the control of NBG stays (BRRD 45f(2)(b)).")
    row("s_int_gap", "Internal MREL of BIL: instruments that NBG must buy (EUR bn)", [f"=MAX(0,B{S.rows['s_int_req']}-B{S.rows['s_int_res']})"], "bn", "NBG issues them outside the group and passes the money to BIL. At group level the instruments cancel.", None, True)
    S.section("Multiple points of entry: BIL stays its own resolution entity")
    S.row("m_b_sur", "BIL alone: surplus (+) or shortfall (-) against its requirement (EUR bn)", [None, "=Calc!E25-Calc!E181"], "bn", "BIL keeps its own resolution plan and its own requirement. Its senior debt counts for itself.")
    S.row("m_n_req", "NBG alone: requirement with the buffer (EUR bn)", [None, "=Calc!D181"], "bn")
    S.row("m_stake", "NBG exposure to BIL: price paid for the stake (EUR bn)", [None, "=Calc!E38"], "bn", "Calc row 38.")
    S.row("m_add_pct", "Add-on of the NBG requirement for the equity stake in BIL (% of the stake)", [None, 1.0], "pct", "Judgment. The SRB policy says that MREL rises for exposures between resolution groups, because the equity of a participation may have to be written off (paragraphs 41 and 42). 0 = no add-on.", [None, "input"])
    S.row("m_add", "Add-on (EUR bn)", [None, f"=C{S.rows['m_stake']}*C{S.rows['m_add_pct']}"], "bn")
    S.row("m_n_res", "NBG alone: resources (EUR bn)", [None, "=Calc!D25"], "bn", "Simplification: NBG resources before the deal. The cash paid for BIL and the goodwill lower the CET1 of NBG: see Calc rows 20 and 74.")
    S.row("m_n_sur", "NBG alone: surplus (+) or shortfall (-) with the add-on (EUR bn)", [None, f"=C{S.rows['m_n_res']}-C{S.rows['m_n_req']}-C{S.rows['m_add']}"], "bn")
    S.row("m_need", "Issuance need of the two groups together (EUR bn); a surplus of one does not cover the shortfall of the other", [None, f"=MAX(0,-C{S.rows['m_b_sur']})+MAX(0,-C{S.rows['m_n_sur']})"], "bn", "", None, True)
    S.section("Selected strategy")
    row("sel_need", "Issuance need at the requirements, selected strategy (EUR bn)", [None, None, f"=IF($B$4=1,B{S.rows['s_need_req']},C{S.rows['m_need']})"], "bn", "Single point of entry: Calc row 192. Multiple points of entry: the sum of the two shortfalls.", None, True)
    row("sel_gap", "Instruments that NBG must buy in BIL, selected strategy (EUR bn)", [None, None, f"=IF($B$4=1,B{S.rows['s_int_gap']},0)"], "bn")
    S.skip()
    br = pack.get("bridge") or {}
    if br:
        bn, bl = br["NBG"], br["BIL"]
        both = [bn, bl]
        g = lambda f: [f(b) for b in both]  # noqa: E731
        bnm = lambda f: [None if f(b) is None else f(b) / 1000.0 for b in both]  # noqa: E731  P3DH amounts are in EUR m: this sheet is in EUR bn
        tl = lambda k: bnm(lambda b: b["tlac"].get(k))  # noqa: E731
        data2, cols2 = ["data", "data"], ("B", "C")
        S.head(["E. Issued debt and MREL-eligible debt (P3DH data: CC2 liabilities, TLAC1, KM2)", "NBG", "BIL", "", "", "", "", "Source and note"])
        row("b_date", "Date of CC2 and TLAC1", g(lambda b: b["date"]), "text", "Newest date that has both templates. The dates of the two banks can differ.", data2)
        row("b_debt", "Debt securities issued (CC2)", bnm(lambda b: b["cc2"].get("debt")), "bn", "P3DH CC2, liabilities. NBG: debt securities in issue.", data2)
        row("b_fv", "Financial liabilities at fair value: structured notes (CC2)", bnm(lambda b: b["cc2"].get("fv")), "bn", "BIL designates its structured notes at fair value.", data2)
        row("b_sub", "Subordinated debts and other borrowed funds (CC2)", bnm(lambda b: b["cc2"].get("sub")), "bn", "NBG holds its Tier 2 inside the debt securities line.", data2)
        row("b_at1", "AT1 (CC2 other equity instruments, else TLAC1 row 2)", bnm(lambda b: b["cc2"].get("at1") or b["tlac"].get("at1")), "bn", "NBG reports its AT1 in equity, outside the liability lines.", data2)
        row("b_issued", "Total issued debt", [f"=SUM({c}{S.rows['b_debt']}:{c}{S.rows['b_at1']})" for c in cols2], "bn", "", None, True)
        row("b_e_at1", "Eligible: AT1 (TLAC1 row 2)", tl("at1"), "bn", "", data2)
        row("b_e_t2", "Eligible: Tier 2, with the part of 1 to 5 years (rows 6 and EU-12c)", bnm(lambda b: (b["tlac"].get("t2") or 0) + (b["tlac"].get("t2_am") or 0)), "bn", "", data2)
        row("b_e_snp", "Eligible: senior non-preferred (rows 12 and EU-12b)", bnm(lambda b: (b["tlac"].get("snp") or 0) + (b["tlac"].get("snp_b") or 0)), "bn", "", data2)
        row("b_e_ns", "Eligible: non-subordinated (row 14, else rows 13 and EU-13a)",
            bnm(lambda b: b["tlac"]["ns14"] if b["tlac"].get("ns14") is not None else (b["tlac"].get("ns13") or 0) + (b["tlac"].get("ns13a") or 0)), "bn", "", data2)
        row("b_elig", "Eligible debt, gross", [f"=SUM({c}{S.rows['b_e_at1']}:{c}{S.rows['b_e_ns']})" for c in cols2], "bn", "", None, True)
        row("b_not", "Not eligible (issued less eligible)", [f"={c}{S.rows['b_issued']}-{c}{S.rows['b_elig']}" for c in cols2], "bn",
            "The accounting lines and the MREL types do not map one to one. Structured notes with a variable principal and notes with less than 1 year left are the usual causes.")
        row("b_share", "Eligible share of issued debt", [f"=IF({c}{S.rows['b_issued']}>0,{c}{S.rows['b_elig']}/{c}{S.rows['b_issued']},0)" for c in cols2], "pct")
        row("b_ded", "Deduction of own holdings of other eligible instruments (TLAC1 row 20)", [x or 0.0 for x in bnm(lambda b: b["tlac"].get("ded"))], "bn", "Negative. It reduces the non-subordinated layer.", data2)
        row("b_net", "Eligible debt after deduction", [f"={c}{S.rows['b_elig']}+{c}{S.rows['b_ded']}" for c in cols2], "bn", "", None, True)
        row("b_cet1", "CET1 (TLAC1 row 1)", tl("cet1"), "bn", "", data2)
        row("b_total", "Own funds and eligible liabilities, computed", [f"={c}{S.rows['b_net']}+{c}{S.rows['b_cet1']}" for c in cols2], "bn", "", None, True)
        row("b_trea", "TREA (TLAC1 row 23)", tl("trea"), "bn", "", data2)
        row("b_ratio", "Own funds and eligible liabilities, % of TREA, computed", [f"=IF({c}{S.rows['b_trea']}>0,{c}{S.rows['b_total']}/{c}{S.rows['b_trea']},0)" for c in cols2], "pct2", "", None, True)
        row("b_rep", "Reported (KM2 row 3, same date)", g(lambda b: b["km2"].get("ratio")), "pct2", "", data2)
        row("b_bp", "Difference (bp)", [f"=({c}{S.rows['b_ratio']}-{c}{S.rows['b_rep']})*10000" for c in cols2], "0.00", "Within 0.5 bp is a match: the bank reports the ratio to 0.01 percentage point.")
        row("b_k_date", "Newest KM2: date", g(lambda b: b["km2_new"].get("date")), "text", "The model (sheet Calc) uses this KM2.", data2)
        row("b_k_tot", "Newest KM2: own funds and eligible liabilities", bnm(lambda b: b["km2_new"].get("total")), "bn", "", data2)
        row("b_k_chk", "Equal to Calc row 25 (1 = yes)", [f"=IF(ABS({c}{S.rows['b_k_tot']}-Calc!{k}25)<0.0005,1,0)" for c, k in zip(cols2, ("D", "E"))], "flag",
            "Own funds and eligible liabilities in the model come from the same KM2.")
        S.skip()
    S.head(["F. Caveats and what is not verified", "Status", "", "", "", "", "", "Source"])
    caveats = [
        ("The SRB decides the strategy (single or multiple point of entry). The bank is heard on MREL and can appeal the MREL decision, not the plan.", "Verified", "SRMR Articles 8, 12 to 12k; SRB resolution planning cycle booklet 2024"),
        ("The law does not define single and multiple point of entry. The SRB has no default (68 single and 16 multiple of 84 resolution entities).", "Verified", "BRRD2 recital 4; SRB MREL dashboard H1 2025"),
        ("An acquired resolution entity normally stops being a resolution entity. Its MREL decision stays valid until the SRB replaces it. There is no grandfathering and no transition rule for M&A.", "Verified", "SRB paper on M&A of 7 Dec 2020; BRRD 13(3)"),
        ("Own funds count for MREL. The senior debt of a subsidiary does not count for the group: only the resolution entity issues eligible liabilities.", "Verified", "CRR 72l, 72b(2)(a); SRB MREL policy, paragraph 111"),
        ("Outside-held AT1 and Tier 2 of a subsidiary enter consolidated own funds up to the part that supports the requirement of the subsidiary (section C gives the share for BIL).", "Verified (formula); the share is a calculation", "CRR Articles 84 to 88"),
        ("Internal MREL of a subsidiary needs instruments issued to and bought by the resolution entity, ranking below its other liabilities. No scaling. No waiver across member states.", "Verified", "BRRD 45f(2) to (4); SRB MREL policy, paragraphs 37 and 127"),
        ("Resolution groups above EUR 100bn: 13.5% of TREA and 5% of the leverage exposure in subordinated resources, with the combined buffer on top. A transition of 3 years applies after crossing.", "Verified", "BRRD 45c(5), 45m(4); CRD Article 128"),
        ("The date of the EUR 100bn test is not in the text.", "Unverified", "SRB policy paragraph 13: prior-year data or later data where needed"),
        ("How often the SRB adds a market confidence charge, and whether it would for the combination.", "Unverified", ""),
        ("The SRB will choose single point of entry for NBG + BIL. The precedent of Eurobank and Hellenic Bank points to it. BIL must stay separable for multiple points of entry.", "Inference", "Eurobank reports; SRB M&A paper"),
        ("The BIL senior debt may be called by BIL after an MREL disqualification event. A change of group may trigger it.", "Unverified", "BIL EMTN base prospectus 2025"),
        ("The add-on for the equity stake under multiple points of entry is a judgment.", "Inference", "SRB MREL policy, paragraphs 41 and 42"),
        ("The reform of crisis management (Directive (EU) 2026/806) applies from May 2028 and changes none of these rules.", "Verified", "Official Journal 20 April 2026"),
    ]
    for text, status, src in caveats:
        S.row(None, text, [status], None, src, None)
    ws.freeze_panes = "B7"
    return S.rows
