"""Writes the issuance workbook: sheets Read me, Profile (liability profile by debt type: each issuer and combined), Issues (one row for each bond), and Listed without data.
The Profile sheet holds formulas (SUMIFS and COUNTIFS) over the sheets Issues and Listed without data. A change in Issues changes the Profile. Labels follow ASD-STE100."""

from __future__ import annotations

import datetime as dt
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .issuance import AT1, COVERED, DEBT_TYPES, PERPETUAL_KEY, SENIOR, SNP, SP, STRUCTURED, T2, bucket_labels

HEAD, HEAD_FONT = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
SECTION, INPUT = PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="FFF2CC")
BLUE, GREY, BOLD = Font(color="0000FF"), Font(color="808080", italic=True), Font(bold=True)
SHEETS = ("Read me", "Profile", "Bridge", "Issues", "Listed without data")

ISSUE_COLUMNS = [  # (header, key, width, number format)
    ("Issuer", "issuer", 8, None), ("Active", "active", 8, None), ("Debt type", "debt_type", 34, None), ("Public / private", "public_private", 14, None), ("Programme", "programme", 11, None),
    ("Basis of the class", "basis", 38, None), ("ISIN", "isin", 15, None), ("Bloomberg security", "security", 34, None), ("Rank (Bloomberg)", "rank", 17, None), ("Basel III", "basel", 17, None),
    ("Collateral", "collateral", 17, None), ("Coupon type", "coupon_type", 14, None), ("Coupon (%)", "coupon", 10, "0.000"), ("Currency", "currency", 9, None), ("Issue date", "issue_date", 12, "yyyy-mm-dd"),
    ("Maturity", "maturity", 12, "yyyy-mm-dd"), ("Next call", "next_call", 12, "yyyy-mm-dd"), ("Perpetual", "perpetual", 10, None), ("Issued (local, m)", "issued", 14, "#,##0.0"),
    ("Outstanding (local, m)", "outstanding", 16, "#,##0.0"), ("Outstanding (EUR m)", "eur", 16, "#,##0.0"), ("Size bucket (EUR m)", "size_bucket", 13, None), ("Years to maturity", "years", 10, "0.0"),
    ("Maturity bucket", "bucket_m", 10, "0"), ("Call-adjusted bucket", "bucket_c", 11, "0"), ("Price, mid", "px", 9, "0.00"), ("Yield to maturity, mid (%)", "ytm", 12, "0.00"),
    ("I-spread (bp)", "i_spread", 10, "0"), ("Z-spread (bp)", "z_spread", 10, "0"), ("Quote date", "quote_date", 12, "yyyy-mm-dd"), ("Series (Bloomberg)", "series", 10, None), ("Name (Bloomberg)", "name", 30, None)]
UN_COLUMNS = [("Issuer", "issuer", 8, None), ("Bloomberg name", "security", 34, None), ("Status", "status", 12, None), ("Meaning", "meaning", 70, None), ("Coupon (%)", "coupon", 10, "0.000"),
              ("Maturity", "maturity", 12, "yyyy-mm-dd"), ("Perpetual", "perpetual", 10, None), ("Maturity bucket", "bucket", 10, "0"), ("Description (Bloomberg)", "description", 44, None)]


def _col(columns: list, key: str) -> str:
    return get_column_letter([c[1] for c in columns].index(key) + 1)


def _table(ws, frame: pd.DataFrame, columns: list, name: str) -> int:
    for j, (head, key, width, fmt) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=j, value=head)
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD, Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = width
        for i, v in enumerate(frame[key].tolist() if key in frame else [None] * len(frame), start=2):
            if isinstance(v, float) and v != v:
                v = None
            if isinstance(v, pd.Timestamp):
                v = v.date()
            cell = ws.cell(row=i, column=j, value=v)
            if fmt:
                cell.number_format = fmt
    last = max(len(frame) + 1, 2)
    t = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "D2"
    ws.row_dimensions[1].height = 32
    return last


def write_profile(ws, data: dict[str, Any], n_issues: int, n_un: int) -> None:
    banks = data["meta"]["banks"]
    first = data["first_year"]
    labels = bucket_labels(first)
    ic = lambda k: f"Issues!${_col(ISSUE_COLUMNS, k)}$2:${_col(ISSUE_COLUMNS, k)}${n_issues}"  # noqa: E731
    uc = lambda k: f"'Listed without data'!${_col(UN_COLUMNS, k)}$2:${_col(UN_COLUMNS, k)}${n_un}"  # noqa: E731
    ws["A1"] = "Liability profile by debt type: each issuer and combined. Amounts in EUR million, active issues only."
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Black: formula. Blue on yellow: input. The formulas read the sheet Issues: change a debt type or a class there and the profile follows."
    ws["A4"] = "Basis of the profile: 1 = contractual maturity, 2 = next call date where a call exists, else maturity"
    ws["B4"] = 1
    ws["B4"].font, ws["B4"].fill = BLUE, INPUT
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    n_year = len(labels)
    last_col = 1 + n_year  # last column of the buckets
    blocks = [(s, s) for s, _ in banks] + ([("Combined", "*")] if len(banks) > 1 else [])
    r0 = 6
    for title, crit in blocks:
        ws.cell(row=r0, column=1, value=title).font = Font(bold=True, size=12)
        ws.cell(row=r0, column=2, value="Issuer filter in the sheet Issues (* = all):")
        c = ws.cell(row=r0, column=5, value=crit)
        c.font, c.fill = BLUE, INPUT
        crit_ref = f"$E${r0}"
        for j in range(1, last_col + 4):
            ws.cell(row=r0, column=j).fill = SECTION if j not in (5,) else INPUT
        rh, rk = r0 + 1, r0 + 2
        heads = ["Debt type (EUR m)"] + [lab for _, lab in labels] + ["Total", "Share of total", "Notes"]
        for j, h in enumerate(heads, start=1):
            c = ws.cell(row=rh, column=j, value=h)
            c.font, c.fill, c.alignment = HEAD_FONT, HEAD, Alignment(horizontal="left" if j == 1 else "center")
        ws.cell(row=rk, column=1, value="Bucket key (used by the formulas)").font = GREY
        for j, (key, _) in enumerate(labels, start=2):
            ws.cell(row=rk, column=j, value=key).font = GREY
        rows = {}
        for i, kind in enumerate(DEBT_TYPES):
            r = rk + 1 + i
            rows[kind] = r
            ws.cell(row=r, column=1, value=kind)
            for j in range(2, last_col + 1):
                cl = get_column_letter(j)
                base = f'{ic("eur")},{ic("issuer")},{crit_ref},{ic("debt_type")},$A{r},{ic("active")},"Y"'
                ws.cell(row=r, column=j, value=f"=IF($B$4=1,SUMIFS({base},{ic('bucket_m')},{cl}${rk}),SUMIFS({base},{ic('bucket_c')},{cl}${rk}))").number_format = "#,##0.0"
            tot = get_column_letter(last_col + 1)
            ws.cell(row=r, column=last_col + 1, value=f"=SUM(B{r}:{get_column_letter(last_col)}{r})").number_format = "#,##0.0"
            ws.cell(row=r, column=last_col + 3, value=f'=COUNTIFS({ic("issuer")},{crit_ref},{ic("debt_type")},$A{r},{ic("active")},"Y")').number_format = "0"
        rt = rk + 1 + len(DEBT_TYPES)
        ws.cell(row=rt, column=1, value="Total").font = BOLD
        for j in range(2, last_col + 2):
            cl = get_column_letter(j)
            c = ws.cell(row=rt, column=j, value=f"=SUM({cl}{rk + 1}:{cl}{rt - 1})")
            c.number_format, c.font = "#,##0.0", BOLD
        c = ws.cell(row=rt, column=last_col + 3, value=f"=SUM({get_column_letter(last_col + 3)}{rk + 1}:{get_column_letter(last_col + 3)}{rt - 1})")
        c.font = BOLD
        tot = get_column_letter(last_col + 1)
        for r in range(rk + 1, rt + 1):
            ws.cell(row=r, column=last_col + 2, value=f"=IF(${tot}${rt}>0,{tot}{r}/${tot}${rt},0)").number_format = "0.0%"
        ru = rt + 1
        ws.cell(row=ru, column=1, value="Listed notes without data (count by maturity; amount unknown)").font = GREY
        for j, (key, _) in enumerate(labels, start=2):
            cl = get_column_letter(j)
            ws.cell(row=ru, column=j, value=f"=COUNTIFS({uc('issuer')},{crit_ref},{uc('bucket')},{cl}${rk})").font = GREY
        ws.cell(row=ru, column=last_col + 1, value=f"=SUM(B{ru}:{get_column_letter(last_col)}{ru})").font = GREY

        rp = ru + 2
        for j, h in enumerate(["Public and private (EUR m)", "Public", "Private", "Total", "Private share", "Public notes", "Private notes"], start=1):
            c = ws.cell(row=rp, column=j, value=h)
            c.font, c.fill, c.alignment = HEAD_FONT, HEAD, Alignment(horizontal="left" if j == 1 else "center")
        for i, kind in enumerate(DEBT_TYPES):
            r = rp + 1 + i
            ws.cell(row=r, column=1, value=kind)
            for j, cls in ((2, "Public"), (3, "Private")):
                ws.cell(row=r, column=j, value=f'=SUMIFS({ic("eur")},{ic("issuer")},{crit_ref},{ic("debt_type")},$A{r},{ic("public_private")},"{cls}",{ic("active")},"Y")').number_format = "#,##0.0"
                ws.cell(row=r, column=j + 4, value=f'=COUNTIFS({ic("issuer")},{crit_ref},{ic("debt_type")},$A{r},{ic("public_private")},"{cls}",{ic("active")},"Y")').number_format = "0"
            ws.cell(row=r, column=4, value=f"=B{r}+C{r}").number_format = "#,##0.0"
            ws.cell(row=r, column=5, value=f"=IF(D{r}>0,C{r}/D{r},0)").number_format = "0.0%"
        rpt = rp + 1 + len(DEBT_TYPES)
        ws.cell(row=rpt, column=1, value="Total").font = BOLD
        for j in (2, 3, 4, 6, 7):
            cl = get_column_letter(j)
            c = ws.cell(row=rpt, column=j, value=f"=SUM({cl}{rp + 1}:{cl}{rpt - 1})")
            c.font, c.number_format = BOLD, "#,##0.0" if j < 6 else "0"
        ws.cell(row=rpt, column=5, value=f"=IF(D{rpt}>0,C{rpt}/D{rpt},0)").number_format = "0.0%"

        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title = f"{title}: liability profile by debt type (EUR m)"
        ch.y_axis.title = "EUR m"
        ch.add_data(Reference(ws, min_col=1, max_col=last_col, min_row=rk + 1, max_row=rt - 1), from_rows=True, titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, max_col=last_col, min_row=rh))
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.width, ch.height = 24, 9
        ws.add_chart(ch, f"{get_column_letter(last_col + 5)}{r0}")
        r0 = max(rpt + 3, r0 + 22)
    ws.column_dimensions["A"].width = 58
    for j in range(2, last_col + 5):
        ws.column_dimensions[get_column_letter(j)].width = 11
    ws.freeze_panes = "B6"


def read_me_lines(data: dict[str, Any]) -> list[tuple[str, bool]]:
    meta, iss, un = data["meta"], data["issues"], data["unresolved"]
    names = ", ".join(s for s, _ in meta["banks"])
    act = iss[iss["active"] == "Y"]
    lines: list[tuple[str, bool]] = [
        (f"Issuance of {names}: the bonds of each issuer in the Bloomberg tables of the P3DH database, with a liability profile by debt type.", True), ("", False),
        ("What this file holds", True),
        ("Profile: the liability profile by debt type, in EUR million, for each issuer and for all issuers combined. Switch B4 chooses contractual maturity or next call.", False),
        ("Bridge: issued debt (P3DH, CC2) and MREL-eligible debt (P3DH, TLAC1) for each issuer, and a check that the own funds and eligible liabilities equal the reported KM2 ratios. Use this sheet for the size and the structure of the debt.", False),
        ("Issues: one row for each bond. Filter the columns. Column C is the debt type, column D is public or private.", False),
        ("Listed without data: live bond names that the Bloomberg search lists but that have no record in the database. They are not in the amounts of the Profile.", False), ("", False),
        ("Data dates", True),
        (f"Built on {meta['today']}. Static data of the bonds: pulled {str(meta['static_from'])[:10]} to {str(meta['static_to'])[:10]}. Latest quote date: {meta['market_to']}. FX rates for EUR: {meta['fx_date']}.", False),
        ("The sheets Profile and Issues show the nominal amounts outstanding that Bloomberg gives on the build date, not accounting carrying amounts. They include notes issued after the last balance sheet.", False),
        ("Do not compare them with a balance sheet without the sheet Bridge: it uses P3DH (accounts and MREL templates) at the date of the last filing.", False), ("", False),
        ("Active", True),
        ("A bond is active when its amount outstanding is above 0 and its maturity is not in the past (or it is a perpetual). The Profile counts active bonds only.", False), ("", False),
        ("Debt types", True),
        ("AT1: Basel III designation Additional Tier 1, or rank Jr Subordinated. Tier 2: designation Tier 2, or rank Subordinated.", False),
        ("Senior non-preferred: rank Sr Non Preferred. Covered bond: collateral type COVERED. Senior preferred: rank Sr Preferred.", False),
        ("Senior unsecured (programme notes): the other senior notes that are in a public MTN programme, with a fixed, floating, or variable coupon, that do not convert.", False),
        ("Structured and private notes: every other senior note. These are notes with a zero or step coupon, notes that can convert into shares or another asset, and notes outside a programme.", False),
        ("Bloomberg has no field for structured products. The class is a rule on the fields above. Check the rows in the sheet Issues before you use them for a decision.", False), ("", False),
        ("Public or private", True),
        ("Public: the note is in a public MTN programme (series EMTN, GMTN, or DMTN), or it is a benchmark outside a programme (EUR 150m or more).", False),
        ("Private: the note is not in a programme and is smaller than EUR 150m. Many small notes in an EMTN programme are private placements. They are public here because the series is EMTN.", False),
        ("Column F of the sheet Issues gives the basis of each class. Column V gives the size bucket, so you can set your own rule.", False), ("", False),
        ("Counts", True),
    ]
    for s, _ in meta["banks"]:
        a = act[act["issuer"] == s]
        pub = a[a["public_private"] == "Public"]
        pri = a[a["public_private"] == "Private"]
        lines.append((f"{s}: {len(a)} active bonds, EUR {a['eur'].sum():,.0f}m. Public {len(pub)} (EUR {pub['eur'].sum():,.0f}m). Private {len(pri)} (EUR {pri['eur'].sum():,.0f}m). "
                      f"Not active: {int((iss[iss['issuer'] == s]['active'] == 'N').sum())}. Listed without data: {int((un['issuer'] == s).sum())}.", False))
    lines += [("", False), ("Gaps", True)]
    for status, text in (("Ambiguous", "Several notes share the name. Bloomberg gives no single record, so the notes cannot be read one by one."),
                         ("Not found", "Bloomberg returned no record for the name."), ("Not loaded", "The name is new since the last pull of static data. A new pull of static data loads it.")):
        for s, _ in meta["banks"]:
            n = int(((un["issuer"] == s) & (un["status"] == status)).sum())
            if n:
                lines.append((f"{s}, {status.lower()}: {n} names. {text}", False))
    lines += [("The amounts of these notes are not in the Profile. The sheet Listed without data shows them by maturity (coupon and maturity come from the name).", False),
              ("To load the new names: run scripts/pull_bloomberg.py bonds with the LEI of the bank. It needs the Bloomberg terminal and unique securities of the monthly budget.", False),
              ("Rebuild this file: dashboard/.venv/Scripts/python scripts/build_issuance_excel.py", False)]
    return lines


def write_workbook(path: str, data: dict[str, Any], bridge: dict[str, dict[str, Any]] | None = None) -> None:
    """bridge: the result of issuance.load_bridge. Without it the workbook has no sheet Bridge."""
    from .issuance_bridge import write_bridge

    wb = Workbook()
    wb.active.title = SHEETS[0]
    names = [n for n in SHEETS[1:] if n != "Bridge" or bridge is not None]
    sheets = {n: wb.create_sheet(n) for n in names}
    n_i = _table(sheets["Issues"], data["issues"], ISSUE_COLUMNS, "IssuesTable")
    n_u = _table(sheets["Listed without data"], data["unresolved"], UN_COLUMNS, "ListedWithoutData")
    write_profile(sheets["Profile"], data, n_i, n_u)
    if bridge is not None:
        write_bridge(sheets["Bridge"], lambda k: _col(ISSUE_COLUMNS, k), n_i, data["meta"]["banks"], bridge, data["issues"])
    ws = wb["Read me"]
    for i, (text, bold) in enumerate(read_me_lines(data), start=1):
        c = ws.cell(row=i, column=1, value=text)
        if bold:
            c.font = BOLD
    ws.column_dimensions["A"].width = 170
    wb.active = 1
    for sheet in wb.worksheets:  # one selected tab only: two selected tabs put Excel in group-edit mode
        sheet.sheet_view.tabSelected = sheet.title == "Profile"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
